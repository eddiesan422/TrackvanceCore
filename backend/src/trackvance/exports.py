"""Business-readable, immutable XLSX report snapshots. Authorization belongs to the API.

All untrusted strings are serialized as OOXML strings, never formulas. Results keep
their original text; only fields with declared numeric/date semantics are typed.
"""

from __future__ import annotations

import io
import json
import re
from collections.abc import Iterable, Mapping, Sequence
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NAVY = "172D4D"
TEAL = "087F8C"
RED = "B42335"
AMBER = "9A6700"
INK = "24364B"
MUTED = "65758B"
NUMBER = "#,##0.00########"
PERCENT = "0.00%"
DATE = "yyyy-mm-dd"
DATETIME = "yyyy-mm-dd hh:mm:ss"
NA = "No disponible"
MAX_EXCEL_ROWS = 1_048_576
STATUS_LABELS = {
    "COMPLETED": "Completada",
    "SUCCESS": "Completada",
    "FAILED": "Error técnico",
    "FAILED_PRECONDITION": "Precondición incumplida",
    "CANCELLED": "Cancelada",
    "QUEUED": "En cola",
    "RUNNING": "En ejecución",
    "PASS": "Cumple",
    "FAIL": "No cumple",
    "WARN": "Advertencia",
    "SKIPPED": "No evaluada",
    "ERROR_TECHNICAL": "Error técnico",
    "ERROR": "Error de negocio",
    "WARNING": "Advertencia",
    "APPROVED": "Aprobado",
    "REJECTED": "Rechazado",
    "APPROVED_WITH_WARNINGS": "Aprobado con advertencias",
    "HEALTHY": "Saludable",
    "ALERT": "Alerta",
    "WITH_FINDINGS": "Con hallazgos",
    "MATCH": "Coincide",
    "VALUE_MISMATCH": "Diferencia de valor",
    "SOURCE_ONLY": "Solo en origen",
    "TARGET_ONLY": "Solo en destino",
    "DUPLICATE_SOURCE": "Duplicado en origen",
    "DUPLICATE_TARGET": "Duplicado en destino",
    "INVALID": "No evaluable",
}
CLASSIFICATIONS = (
    "MATCH",
    "VALUE_MISMATCH",
    "SOURCE_ONLY",
    "TARGET_ONLY",
    "DUPLICATE_SOURCE",
    "DUPLICATE_TARGET",
    "INVALID",
)
_SECRET_KEY = re.compile(
    r"password|passwd|secret|token|credential|authorization|cookie|private[_-]?key", re.IGNORECASE
)


def export_filename(module: str, run_id: str) -> str:
    """An ASCII basename safe for Content-Disposition, regardless of caller input."""
    safe_module = re.sub(r"[^a-z0-9_-]", "_", module.lower())[:24].strip("_-") or "run"
    safe_id = re.sub(r"[^a-zA-Z0-9_-]", "_", str(run_id))[:80].strip("_-") or "run"
    return f"trackvance_{safe_module}_{safe_id}.xlsx"


def _mapping(value: Any) -> Mapping[str, Any]:
    return value if isinstance(value, Mapping) else {}


def _redact(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(k): _redact(v) for k, v in value.items() if not _SECRET_KEY.search(str(k))}
    if isinstance(value, (list, tuple)):
        return [_redact(v) for v in value]
    return value


def _text(value: Any) -> str:
    if value is None:
        return NA
    if isinstance(value, (Mapping, list, tuple)):
        return json.dumps(_redact(value), ensure_ascii=False, sort_keys=True, default=str)
    return str(value)


def _label(code: Any) -> str:
    value = str(code or NA)
    return f"{STATUS_LABELS[value]} ({value})" if value in STATUS_LABELS else value


def _decimal(value: Any) -> Any:
    """Only called for known numeric metrics/comparison fields, never identifiers."""
    if value is None or isinstance(value, bool):
        return value
    try:
        result = Decimal(str(value))
        if not result.is_finite():
            return str(value)
        # Excel stores at most 15 significant digits; retain evidence beyond this as text.
        if len(result.as_tuple().digits) > 15:
            return str(value)
        return result
    except (InvalidOperation, ValueError):
        return value


def _timestamp(value: Any) -> Any:
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return value
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(UTC).replace(tzinfo=None)
    return value


def _ratio(value: Any, scale: int = 100) -> Any:
    number = _decimal(value)
    return number / scale if isinstance(number, Decimal) else number


def _write(ws: Worksheet, row: int, column: int, value: Any, fmt: str | None = None) -> Cell:
    cell = ws.cell(row, column)
    assert isinstance(cell, Cell)
    if value is None:
        value = NA
    if isinstance(value, (Mapping, list, tuple)):
        value = _text(value)
    if isinstance(value, str):
        # Excel limits text to 32767 characters. Fail explicitly instead of hiding evidence.
        if len(value) > 32767:
            raise ValueError("Un valor supera el límite de texto de Excel (32767 caracteres).")
        # XML 1.0 cannot encode these controls. Escape visibly; retain tabs/newlines verbatim.
        value = ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\u{ord(m[0]):04x}", value)
        cell.value = value
        cell.data_type = "s"
        cell.number_format = "@"
    elif isinstance(value, datetime):
        cell.value = _timestamp(value)
        cell.number_format = DATETIME
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = DATE
    else:
        cell.value = value
        cell.number_format = "#,##0" if isinstance(value, int) else NUMBER
    if fmt and not isinstance(value, str):
        cell.number_format = fmt
    cell.font = Font(name="Arial", size=10, color=INK)
    cell.alignment = Alignment(
        vertical="center",
        horizontal="left" if isinstance(value, str) else "right",
        wrap_text=True,
        indent=1,
    )
    return cell


def _sheet(
    book: Workbook, name: str, title: str, subtitle: str, widths: Sequence[int]
) -> Worksheet:
    ws = book.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = NAVY if name == "Resumen" else TEAL
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.sheet_format.defaultRowHeight = 25
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 29
    _write(ws, 2, 1, f"Trackvance Core · {title}").font = Font(
        name="Arial", size=16, bold=True, color=NAVY
    )
    ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    _write(ws, 3, 1, subtitle).font = Font(name="Arial", size=10, color=MUTED)
    ws.cell(3, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
        ws.cell(3, index).border = Border(bottom=Side(style="thin", color=TEAL))
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    return ws


def _table(
    ws: Worksheet,
    name: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    start_row: int = 5,
    start_col: int = 1,
    formats: Mapping[int, str] | None = None,
) -> int:
    for index, header in enumerate(headers, start_col):
        cell = _write(ws, start_row, index, header)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[start_row].height = 30
    end = start_row
    for end, values in enumerate(rows, start_row + 1):
        if end > MAX_EXCEL_ROWS:
            raise ValueError("El informe supera el máximo de filas de Excel. Use un filtro menor.")
        needed_height = 25
        for index, value in enumerate(values, start_col):
            _write(ws, end, index, value, (formats or {}).get(index - start_col))
            width = ws.column_dimensions[get_column_letter(index)].width or 20
            lines = sum(
                max(1, (len(line) + int(width) - 1) // int(width))
                for line in _text(value).splitlines()
            )
            needed_height = max(needed_height, min(409, lines * 14 + 9))
        ws.row_dimensions[end].height = needed_height
    if end == start_row:
        end += 1
        _write(ws, end, start_col, "Sin registros")
    ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col + len(headers) - 1)}{end}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    if not ws.auto_filter.ref:
        ws.auto_filter.ref = ref
    return end


def _conditional(ws: Worksheet, column: str, first: int, last: int, width: int) -> None:
    """Only fixed, application-authored formulas enter conditional-format rules."""
    groups = (
        (("MATCH", "PASS", "HEALTHY", "APPROVED"), TEAL, "E9F6F3"),
        (("VALUE_MISMATCH", "INVALID", "FAIL", "ERROR", "REJECTED", "ALERT"), RED, "FDECEE"),
        (
            (
                "SOURCE_ONLY",
                "TARGET_ONLY",
                "DUPLICATE_SOURCE",
                "DUPLICATE_TARGET",
                "WARN",
                "WARNING",
            ),
            AMBER,
            "FFF5DD",
        ),
    )
    for codes, color, fill in groups:
        formula = (
            "OR("
            + ",".join(f'ISNUMBER(SEARCH("({code})",${column}{first}))' for code in codes)
            + ")"
        )
        ws.conditional_formatting.add(
            f"A{first}:{get_column_letter(width)}{max(first, last)}",
            FormulaRule(
                formula=[formula],
                fill=PatternFill("solid", fgColor=fill),
                font=Font(color=color),
                stopIfTrue=True,
            ),
        )


def _summary(
    book: Workbook,
    module: str,
    run: Mapping[str, Any],
    configuration: Mapping[str, Any],
    inputs: Sequence[Mapping[str, Any]],
    manifest: Mapping[str, Any],
) -> tuple[Worksheet, Mapping[str, Any], Mapping[str, Any], int]:
    titles = {"intake": "Intake", "recon": "ReconOps", "sentinel": "Sentinel"}
    ws = _sheet(
        book,
        "Resumen",
        titles[module],
        "Informe de ejecución · Fechas en UTC",
        [29, 47, 24, 30, 35],
    )
    config = _mapping(_mapping(manifest.get("configuration")).get("config", configuration.get("config", configuration)))
    metrics = _mapping(run.get("metrics", manifest.get("metrics")))
    processing = _mapping(manifest.get("processing", run.get("execution_plan")))
    input_version = inputs[0] if inputs else {}
    kind = {"intake": "Contrato", "recon": "Control", "sentinel": "Monitor"}[module]
    identity = [
        (kind, configuration.get("name", run.get("name"))),
        (f"Versión del {kind.lower()}", configuration.get("version")),
        (
            "Dataset" if module != "recon" else "Dataset origen",
            input_version.get("dataset_name", run.get("dataset_name")),
        ),
        ("Versión del dataset", input_version.get("version")),
        ("DatasetVersion", input_version.get("id", run.get("dataset_version_id"))),
        ("Run ID", run.get("id", run.get("run_id"))),
        ("Inicio (UTC)", _timestamp(run.get("started_at", manifest.get("started_at")))),
        ("Finalización (UTC)", _timestamp(run.get("finished_at", manifest.get("finished_at")))),
        ("Estado técnico", _label(run.get("status"))),
        (
            "Decisión" if module == "intake" else "Resultado de negocio",
            _label(run.get("decision", metrics.get("decision"))),
        ),
        ("Motor", processing.get("engine", processing.get("engine_family"))),
        (
            "Versión del motor",
            processing.get("engine_runtime_version", processing.get("engine_version")),
        ),
    ]
    if module == "recon":
        target = inputs[1] if len(inputs) > 1 else {}
        identity.extend(
            [
                (
                    "Dataset destino",
                    target.get("dataset_name", configuration.get("target_dataset_name")),
                ),
                ("Versión destino", target.get("version")),
                ("DatasetVersion destino", target.get("id")),
                ("Claves", ", ".join(config.get("key_columns", []))),
            ]
        )
    end = _table(ws, "Identidad", ["Ejecución", "Valor"], identity)
    if module == "intake":
        values = [
            ("Filas recibidas", metrics.get("total_rows")),
            ("Filas válidas", metrics.get("valid_rows")),
            ("Filas con error", metrics.get("error_rows")),
            ("Filas con advertencia", metrics.get("warning_rows")),
            ("Aceptación", _ratio(metrics.get("acceptance_rate"))),
            ("Máxima tasa de error", _ratio(config.get("max_error_rate"), 1)),
        ]
    elif module == "recon":
        counts = _mapping(metrics.get("counts"))
        values = [
            ("Filas origen", metrics.get("source_rows")),
            ("Filas destino", metrics.get("target_rows")),
        ]
        values.extend((_label(code), counts.get(code)) for code in CLASSIFICATIONS)
        values.append(("Tasa de coincidencia", _ratio(metrics.get("match_rate"))))
    else:
        values = [
            ("Filas observadas", metrics.get("row_count")),
            ("Controles evaluados", metrics.get("total_checks")),
            ("Controles fallidos", metrics.get("failed_checks")),
            ("Salud", _ratio(metrics.get("health_score"))),
        ]
    metric_end = _table(ws, "Metricas", ["Indicador", "Resultado"], values, start_col=4)
    for row in range(6, metric_end + 1):
        if ws.cell(row, 4).value in {
            "Aceptación",
            "Máxima tasa de error",
            "Tasa de coincidencia",
            "Salud",
        }:
            ws.cell(row, 5).number_format = PERCENT
            ws.cell(row, 5).font = Font(name="Arial", size=10, bold=True, color=TEAL)
    return ws, config, metrics, max(end, metric_end) + 3


def _intake(
    book: Workbook,
    ws: Worksheet,
    config: Mapping[str, Any],
    metrics: Mapping[str, Any],
    results: list[Mapping[str, Any]],
    row: int,
) -> None:
    rule_rows = [
        (
            r.get("code", r.get("rule_code")),
            r.get("column"),
            _label(r.get("status")),
            r.get("evaluated_count"),
            r.get("failed_count"),
        )
        for r in metrics.get("rules", [])
    ]
    _table(
        ws,
        "ResumenReglas",
        ["Regla", "Columna", "Estado", "Evaluadas", "Incumplimientos"],
        rule_rows,
        start_row=row,
    )
    errors = _sheet(
        book,
        "Errores",
        "Errores de Intake",
        "Una línea puede incumplir varias reglas; los conteos de filas no se suman por regla.",
        [20, 26, 24, 32, 22, 58, 28],
    )
    end = _table(
        errors,
        "ErroresIntake",
        [
            "Registro de la versión" if metrics.get("source_row_numbering") == "RECORD_NUMBER" else "Línea del archivo",
            "Regla",
            "Columna",
            "Valor recibido",
            "Severidad",
            "Mensaje",
            "Clasificación",
        ],
        (
            (
                r.get("original_row_number"),
                r.get("rule_code"),
                r.get("column"),
                r.get("received_value"),
                _label(r.get("severity")),
                r.get("message"),
                _label(r.get("classification")),
            )
            for r in results
        ),
    )
    _conditional(errors, "G", 6, end, 7)
    rules = _sheet(
        book,
        "Reglas",
        "Reglas del contrato",
        "Configuración efectiva de la versión utilizada en esta ejecución.",
        [37, 90],
    )
    rows = _flatten(_redact(config))
    _table(rules, "ConfiguracionIntake", ["Parámetro", "Configuración efectiva"], rows)


def _flatten(config: Mapping[str, Any], prefix: str = "") -> list[tuple[str, Any]]:
    output: list[tuple[str, Any]] = []
    for key, value in config.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, Mapping):
            output.extend(_flatten(value, name))
        elif isinstance(value, list) and value and all(isinstance(v, Mapping) for v in value):
            for index, item in enumerate(value, 1):
                output.extend(_flatten(item, f"{name}[{index}]"))
        else:
            output.append((name, value))
    return output


def _recon(
    book: Workbook,
    ws: Worksheet,
    config: Mapping[str, Any],
    results: list[Mapping[str, Any]],
    row: int,
    metrics: Mapping[str, Any],
) -> None:
    normalization = _mapping(config.get("key_normalization"))
    options = [("Claves", ", ".join(config.get("key_columns", []))),
               ("Normalización de claves", _normalization_description(normalization))]
    aggregation = _mapping(config.get("aggregation"))
    if aggregation:
        options.append(("Agregación 1:N", f"Lado {aggregation.get('side')}; {aggregation.get('operation')} de {aggregation.get('column', 'registros')} → {aggregation.get('output_column')}"))
    end = _table(
        ws,
        "OpcionesConciliacion",
        ["Configuración", "Valor efectivo"],
        options,
        start_row=row,
    )
    comparisons = config.get("comparison_rules", [])
    if not comparisons and config.get("amount_column"):
        comparisons = [{"type": "numeric_tolerance", "source_column": config['amount_column'],
                        "target_column": config['amount_column'], "parameters": {"abs": config.get('tolerance')}}]
    _table(ws, "Comparaciones", ["Campo origen", "Campo destino", "Comparación", "Tolerancia / política", "Normalización de valores"],
           [(c.get("source_column"), c.get("target_column"), c.get("code", str(c.get("type", "")).upper()),
             _comparison_policy(c), _normalization_description(_mapping(c.get("parameters")).get("normalization", {})))
            for c in comparisons], start_row=end + 3)
    headers = [
        "Key",
        "Clasificación",
        "Valor origen",
        "Valor destino",
        "Diferencia",
        "Tolerancia",
        "Registro origen" if metrics.get("source_row_numbering") == "RECORD_NUMBER" else "Línea origen",
        "Registro destino" if metrics.get("target_row_numbering") == "RECORD_NUMBER" else "Línea destino",
        "Mensaje",
        "Detalle de comparaciones",
    ]
    widths = [26, 39, 26, 26, 23, 23, 17, 17, 60, 70]
    for name, title, rows in [
        ("Resultados", "Resultados de ReconOps", results),
        (
            "Hallazgos",
            "Hallazgos de ReconOps",
            sorted(
                (r for r in results if r.get("classification") != "MATCH"),
                key=lambda r: (str(r.get("classification")), str(r.get("key"))),
            ),
        ),
    ]:
        sheet = _sheet(
            book,
            name,
            title,
            "Resultados agrupados por clasificación."
            if name == "Hallazgos"
            else "La clasificación de negocio es independiente del estado técnico del run.",
            widths,
        )
        end = _table(
            sheet,
            f"Recon{name}",
            headers,
            (
                (
                    r.get("key"),
                    _label(r.get("classification")),
                    _recon_value(r, config, "source"),
                    _recon_value(r, config, "target"),
                    _decimal(r.get("difference")),
                    _decimal(r.get("tolerance")),
                    r.get("source_row"),
                    r.get("target_row"),
                    r.get("message"),
                    _comparison_details(r),
                )
                for r in rows
            ),
        )
        sheet.freeze_panes = "C6"
        _conditional(sheet, "B", 6, end, 10)


def _recon_value(result: Mapping[str, Any], config: Mapping[str, Any], side: str) -> Any:
    value = result.get(f"{side}_value")
    comparisons = result.get("comparisons", [])
    code = comparisons[0].get("code") if comparisons else ""
    if code in {"NUMERIC_TOLERANCE", "AGGREGATE_COMPARE"} or (
        not comparisons and "amount_column" in config
    ):
        return _decimal(value)
    if code == "DATE_TOLERANCE":
        return _timestamp(value)
    return value


def _comparison_details(result: Mapping[str, Any]) -> Any:
    comparisons = result.get("comparisons", [])
    if not comparisons:
        return result.get("field_results", result.get("comparison_details"))
    return "\n".join(
        f"{c.get('source_column', '')} / {c.get('target_column', '')}: "
        f"{_text(c.get('source_value'))} / {_text(c.get('target_value'))}. "
        f"{_label(c.get('status', 'PASS' if c.get('passed') else 'FAIL' if 'passed' in c else None))}; {c.get('code', '')}. {c.get('message', '')}"
        for c in comparisons
    )


def _normalization_description(policy: Mapping[str, Any]) -> str:
    return f"Espacios: {'TRIM' if policy.get('trim') else 'NONE'}; case: {policy.get('case', 'NONE')}; Unicode: {policy.get('unicode_normalization', 'NONE')}"


def _comparison_policy(comparison: Mapping[str, Any]) -> str:
    kind, params = comparison.get("type"), _mapping(comparison.get("parameters"))
    nulls = "; nulos iguales" if params.get("equal_nulls") else "; nulos no coinciden"
    if kind == "numeric_tolerance":
        text = f"Absoluta: {params.get('abs', 0)}"
        if params.get("percent") is not None:
            text += f"; {params['percent']}% sobre {params.get('denominator', 'SOURCE')}; cero: EXACT_ONLY"
        return text + nulls
    if kind == "date_tolerance":
        return (f"{params['hours']} horas" if "hours" in params else f"{params.get('days', 0)} días") + "; UTC" + nulls
    return "Igualdad exacta" + nulls


def _sentinel(book: Workbook, metrics: Mapping[str, Any], results: list[Mapping[str, Any]]) -> None:
    checks = results or metrics.get("checks", [])
    headers = ["Control", "Code", "Estado", "Valor observado", "Valor esperado", "Detalle"]
    for name, title, rows in [
        ("Controles", "Controles de Sentinel", checks),
        ("Hallazgos", "Hallazgos de Sentinel", [r for r in checks if r.get("status") == "FAIL"]),
    ]:
        sheet = _sheet(
            book,
            name,
            title,
            "Los controles fallidos describen problemas de datos; no errores técnicos.",
            [29, 35, 22, 33, 55, 68],
        )
        end = _table(
            sheet,
            f"Sentinel{name}",
            headers,
            (
                (
                    r.get("name", r.get("code")),
                    r.get("code"),
                    _label(r.get("status")),
                    r.get("actual", r.get("observed_value")),
                    r.get("expected"),
                    r.get("message", r.get("detail")),
                )
                for r in rows
            ),
        )
        _conditional(sheet, "C", 6, end, 6)


def _trace(
    book: Workbook,
    run: Mapping[str, Any],
    configuration: Mapping[str, Any],
    inputs: Sequence[Mapping[str, Any]],
    manifest: Mapping[str, Any],
) -> None:
    ws = _sheet(
        book,
        "Trazabilidad",
        "Trazabilidad de la ejecución",
        "Evidencia del run original. Los datos no disponibles en manifests históricos se indican expresamente.",
        [31, 36, 91, 27, 24],
    )
    config_evidence = _mapping(manifest.get("configuration"))
    processing = _mapping(manifest.get("processing", run.get("execution_plan")))
    rows = [
        ("Ejecución", "run_id", run.get("id", run.get("run_id"))),
        ("Manifest", "schema_version", manifest.get("schema_version")),
        ("Actor", "initiated_by", manifest.get("initiated_by")),
        ("Configuración", "ID", configuration.get("id", config_evidence.get("id"))),
        ("Configuración", "Versión", configuration.get("version", config_evidence.get("version"))),
        ("Configuración", "config_hash", config_evidence.get("config_hash")),
        ("Procesamiento", "Motor", processing.get("engine", processing.get("engine_family"))),
        (
            "Procesamiento",
            "Versión del motor",
            processing.get("engine_runtime_version", processing.get("engine_version")),
        ),
        ("Procesamiento", "Versión Trackvance", manifest.get("engine_version")),
        (
            "Procesamiento",
            "Política de claves",
            _mapping(config_evidence.get("config", configuration.get("config"))).get("key_normalization"),
        ),
    ]
    manifest_inputs = manifest.get("inputs", [])
    for index, item in enumerate(inputs):
        evidence: Mapping[str, Any] = next(
            (v for v in manifest_inputs if v.get("dataset_version_id") == item.get("id")), {}
        )
        section = "Input origen" if index == 0 else "Input destino"
        for label, value in [
            ("Dataset", item.get("dataset_name")),
            ("Dataset ID", item.get("dataset_id")),
            ("DatasetVersion ID", item.get("id")),
            ("Versión", item.get("version")),
            ("Fuente", item.get("source_type")),
            (
                "Input SHA-256",
                evidence.get("artifact_sha256", evidence.get("sha256", item.get("sha256"))),
            ),
            ("Schema hash", evidence.get("schema_hash", item.get("schema_hash"))),
            ("Original artifact ID", item.get("original_artifact_id")),
            ("Canonical artifact ID", item.get("canonical_artifact_id")),
            ("Canonical SHA-256", evidence.get("canonical_sha256")),
            ("Numeración de registros", evidence.get("row_numbering", _mapping(item.get("profile")).get("row_numbering"))),
        ]:
            rows.append((section, label, value))
    end = _table(ws, "EvidenciaRun", ["Recurso", "Referencia", "Valor"], rows)
    config_rows = _flatten(_redact(_mapping(config_evidence.get("config", configuration.get("config")))))
    end = _table(ws, "ConfiguracionEvidencia", ["Parámetro", "Configuración efectiva"], config_rows, start_row=end + 3)
    _table(
        ws,
        "ArtefactosResultado",
        ["Artifact ID", "Kind", "Nombre", "SHA-256", "Tamaño (bytes)"],
        (
            (
                a.get("artifact_id"),
                a.get("kind"),
                a.get("name"),
                a.get("sha256"),
                a.get("size_bytes"),
            )
            for a in manifest.get("result_artifacts", [])
        ),
        start_row=end + 3,
    )
    # Artifact hashes need a readable width; metadata values above use the wide C column.
    ws.column_dimensions["D"].width = 73


def build_run_workbook(
    run: Mapping[str, Any],
    configuration: Mapping[str, Any],
    inputs: Sequence[Mapping[str, Any]],
    manifest: Mapping[str, Any],
    results: Iterable[Mapping[str, Any]],
) -> bytes:
    """Build one four-sheet local XLSX snapshot; never mutate run/configuration evidence."""
    module = str(run.get("module", manifest.get("module", ""))).lower()
    if module not in {"intake", "recon", "sentinel"}:
        raise ValueError(f"Módulo no soportado para exportación: {module}")
    book = Workbook()
    book.remove(book.worksheets[0])
    book.properties.creator = "Trackvance Core"
    book.properties.title = f"Trackvance Core - {module}"
    book.properties.description = "Informe de resultados y evidencia inmutable de ejecución."
    records = list(results)
    ws, config, metrics, row = _summary(book, module, run, configuration, inputs, manifest)
    if module == "intake":
        _intake(book, ws, config, metrics, records, row)
    elif module == "recon":
        _recon(book, ws, config, records, row, metrics)
    else:
        _sentinel(book, metrics, records)
    _trace(book, run, configuration, inputs, manifest)
    for sheet in book.worksheets:
        sheet.print_area = sheet.calculate_dimension()
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
