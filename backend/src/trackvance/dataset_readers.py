"""Pluggable local dataset readers.

Every connector returns the same observed Polars frame before profiling and
canonical Parquet persistence.  The quality engine therefore has no knowledge
of file formats and future database/object-store connectors can implement the
same ``DatasetReader`` contract.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
import zipfile
from abc import ABC, abstractmethod
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, ClassVar, Protocol, runtime_checkable
from xml.etree.ElementTree import ParseError

import duckdb
import polars as pl
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import MAX_ROWS, MAX_UPLOAD_BYTES
from .processing import ProcessingError

MAX_COLUMNS = 100
INSPECTION_ROWS = 100
TEXT_SAMPLE_BYTES = 64 * 1024
EXCEL_EMPTY_ROW_SCAN_LIMIT = 1_000
MAX_XLSX_EXPANDED_BYTES = max(100 * 1024 * 1024, MAX_UPLOAD_BYTES * 20)
MAX_PARQUET_EXPANDED_BYTES = max(128 * 1024 * 1024, MAX_UPLOAD_BYTES * 20)
MAX_CELL_TEXT_BYTES = 64 * 1024


class UnsupportedDatasetFormat(ProcessingError):
    """Raised when no registered reader can safely identify a source."""


@dataclass(frozen=True)
class ReaderOptions:
    """Format-specific options kept outside the quality configuration."""

    sheet_name: str | None = None
    delimiter: str | None = None

    @classmethod
    def from_mapping(cls, value: Mapping[str, Any] | None) -> ReaderOptions:
        raw = dict(value or {})
        unknown = set(raw) - {"sheet_name", "delimiter"}
        if unknown:
            raise ProcessingError(
                f"Opciones de lectura no soportadas: {', '.join(sorted(unknown))}."
            )
        for key in ("sheet_name", "delimiter"):
            if raw.get(key) is not None and not isinstance(raw[key], str):
                raise ProcessingError(f"La opción {key} debe ser texto.")
        sheet_name = raw.get("sheet_name") or None
        delimiter = _normalize_delimiter(raw.get("delimiter"))
        return cls(sheet_name=sheet_name, delimiter=delimiter)

    def as_dict(self) -> dict[str, str]:
        return {
            key: value
            for key, value in {
                "sheet_name": self.sheet_name,
                "delimiter": self.delimiter,
            }.items()
            if value is not None
        }


@dataclass
class DatasetReadResult:
    """Normalized output consumed by profiling and canonical persistence."""

    frame: pl.DataFrame
    source_format: str
    format_label: str
    media_type: str
    row_numbering: str
    native_schema: dict[str, str] = field(default_factory=dict)
    sheets: list[str] = field(default_factory=list)
    selected_sheet: str | None = None
    detected_delimiter: str | None = None
    row_count: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def reader_metadata(self) -> dict[str, Any]:
        result = {
            key: value
            for key, value in {
                "sheet_name": self.selected_sheet,
                "delimiter": self.detected_delimiter,
                "native_schema": self.native_schema or None,
            }.items()
            if value is not None
        }
        result.update(self.metadata)
        return result


@runtime_checkable
class DatasetSource(Protocol):
    """Acquisition port that produces the quality engine's common dataset model.

    A source may delegate to a file reader, query a database, download from
    object storage or call an API.  Credentials and transport stay in the
    adapter; downstream profiling and rules only receive ``DatasetReadResult``.
    """

    source_kind: str

    def read(
        self, options: ReaderOptions | Mapping[str, Any] | None = None, *, inspect: bool = False
    ) -> DatasetReadResult: ...


class DatasetReader(ABC):
    """Source adapter contract. Implementations must never apply business transforms."""

    format: ClassVar[str]
    format_label: ClassVar[str]
    media_type: ClassVar[str]
    extensions: ClassVar[frozenset[str]]

    @abstractmethod
    def read(
        self, path: Path, options: ReaderOptions, *, limit: int | None = None
    ) -> DatasetReadResult:
        """Read observed source values into the common Polars representation."""

    def inspect(self, path: Path, options: ReaderOptions) -> DatasetReadResult:
        return self.read(path, options, limit=INSPECTION_ROWS)


def _normalize_delimiter(value: str | None) -> str | None:
    if value is None or value == "":
        return None
    aliases = {"\\t": "\t", "TAB": "\t", "tab": "\t"}
    value = aliases.get(value, value)
    if (
        len(value) != 1
        or value in {'"', "'", "\r", "\n", "\\"}
        or value.isalnum()
        or len(value.encode("utf-8")) != 1
    ):
        raise ProcessingError(
            "El delimitador debe ser un único carácter ASCII no alfanumérico."
        )
    return value


def _text_sample(path: Path) -> str:
    try:
        with path.open(encoding="utf-8-sig", newline="") as source:
            sample = source.read(TEXT_SAMPLE_BYTES)
    except UnicodeDecodeError as exc:
        raise ProcessingError("El archivo de texto debe usar codificación UTF-8.") from exc
    if "\x00" in sample:
        raise ProcessingError("El archivo de texto contiene bytes nulos no válidos.")
    return sample


def _detect_delimiter(sample: str, *, csv_legacy_default: bool) -> str:
    nonempty = [line for line in sample.splitlines() if line.strip()]
    if csv_legacy_default:
        first_line = nonempty[0] if nonempty else ""
        return ";" if first_line.count(";") > first_line.count(",") else ","
    candidate = "\n".join(nonempty[:20])
    try:
        return csv.Sniffer().sniff(candidate, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {item: (nonempty[0].count(item) if nonempty else 0) for item in ",;\t|"}
        delimiter, count = max(counts.items(), key=lambda item: item[1])
        if count:
            return delimiter
        raise ProcessingError(
            "No fue posible detectar el delimitador del TXT; selecciónalo explícitamente."
        ) from None


def _headers_from_text(sample: str, delimiter: str) -> list[str]:
    try:
        raw = next(csv.reader(io.StringIO(sample), delimiter=delimiter), [])
    except csv.Error as exc:
        raise ProcessingError("No fue posible leer los encabezados del archivo delimitado.") from exc
    return [str(value) for value in raw]


def _validate_headers(headers: list[str]) -> None:
    if not headers or len(headers) != len(set(headers)):
        raise ProcessingError(
            "El archivo necesita encabezados únicos; hay nombres de columnas duplicados."
        )
    if len(headers) > MAX_COLUMNS:
        raise ProcessingError(f"Este prototipo admite hasta {MAX_COLUMNS} columnas.")
    if any(not name.strip() or name.startswith("__tv_") for name in headers):
        raise ProcessingError(
            "El archivo necesita encabezados válidos, sin prefijo reservado __tv_."
        )


def _check_frame(frame: pl.DataFrame, *, enforce_row_limit: bool) -> pl.DataFrame:
    _validate_headers(frame.columns)
    if enforce_row_limit and frame.height > MAX_ROWS:
        raise ProcessingError(f"Este prototipo admite hasta {MAX_ROWS:,} filas por archivo.")
    return frame


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        rendered = format(value, "f")
        if len(rendered.encode("utf-8")) > MAX_CELL_TEXT_BYTES:
            raise ProcessingError("Un valor numérico supera el tamaño máximo permitido.")
        return rendered
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        return _cell_text(Decimal(repr(value)))
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)
    return str(value)


def _native_column_type(values: list[Any]) -> str:
    observed = [value for value in values if value is not None]
    if not observed:
        return "String"
    if all(isinstance(value, bool) for value in observed):
        return "Boolean"
    if all(isinstance(value, datetime) for value in observed):
        return "Datetime"
    if all(isinstance(value, date) and not isinstance(value, datetime) for value in observed):
        return "Date"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in observed):
        return "Int64"
    if all(
        isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
        for value in observed
    ):
        return "Decimal"
    return "String"


def _frame_from_rows(headers: list[str], rows: list[list[Any]]) -> pl.DataFrame:
    _validate_headers(headers)
    normalized = [[_cell_text(value) for value in row] for row in rows]
    return pl.DataFrame(
        normalized,
        schema={name: pl.String for name in headers},
        orient="row",
    )


def _string_frame(frame: pl.DataFrame) -> pl.DataFrame:
    """Preserve observed scalar values using the quality engine's common string model."""
    return _frame_from_rows(frame.columns, [list(row) for row in frame.rows()])


class DelimitedDatasetReader(DatasetReader):
    csv_legacy_default: ClassVar[bool] = False

    def read(
        self, path: Path, options: ReaderOptions, *, limit: int | None = None
    ) -> DatasetReadResult:
        sample = _text_sample(path)
        delimiter = options.delimiter or _detect_delimiter(
            sample, csv_legacy_default=self.csv_legacy_default
        )
        _validate_headers(_headers_from_text(sample, delimiter))
        row_limit = (MAX_ROWS + 1) if limit is None else limit
        try:
            frame = pl.read_csv(
                path,
                separator=delimiter,
                infer_schema=False,
                encoding="utf8",
                n_rows=row_limit,
            )
        except (pl.exceptions.PolarsError, IndexError) as exc:
            raise ProcessingError(
                "No se pudo leer el archivo delimitado UTF-8. Revisa encabezados y delimitador."
            ) from exc
        _check_frame(frame, enforce_row_limit=limit is None)
        return DatasetReadResult(
            frame=frame,
            source_format=self.format,
            format_label=self.format_label,
            media_type=self.media_type,
            row_numbering="PHYSICAL_LINE",
            native_schema={name: "String" for name in frame.columns},
            detected_delimiter=delimiter,
            row_count=frame.height if limit is None else None,
        )


class CsvDatasetReader(DelimitedDatasetReader):
    format = "CSV"
    format_label = "CSV"
    media_type = "text/csv"
    extensions = frozenset({".csv"})
    csv_legacy_default = True


class TxtDatasetReader(DelimitedDatasetReader):
    format = "TXT"
    format_label = "TXT delimitado"
    media_type = "text/plain"
    extensions = frozenset({".txt", ".tsv"})


class ExcelDatasetReader(DatasetReader):
    format = "XLSX"
    format_label = "Excel XLSX"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extensions = frozenset({".xlsx"})

    @staticmethod
    def _validate_package(path: Path) -> None:
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                    raise ProcessingError("El archivo no contiene un libro Excel XLSX válido.")
                if "xl/vbaProject.bin" in names:
                    raise ProcessingError("Los libros Excel con macros no están permitidos.")
                if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_EXPANDED_BYTES:
                    raise ProcessingError("El contenido expandido del libro Excel supera el límite.")
        except zipfile.BadZipFile as exc:
            raise ProcessingError("No fue posible leer el libro Excel XLSX.") from exc

    @staticmethod
    def _cell_value(cell: Any) -> Any:
        value = cell.value
        if isinstance(value, datetime) and cell.is_date:
            number_format = str(cell.number_format or "")
            if not re.search(r"[hHsS]", number_format):
                return value.date()
        return value

    @classmethod
    def _sheet_frame(
        cls, worksheet: Any, *, limit: int | None
    ) -> tuple[pl.DataFrame, dict[str, str]] | None:
        headers: list[str] | None = None
        raw_columns: list[list[Any]] = []
        rows: list[list[Any]] = []
        empty_scans = 0
        row_limit = MAX_ROWS + 1 if limit is None else limit
        for cells in worksheet.iter_rows():
            values = [cls._cell_value(cell) for cell in cells]
            if headers is None:
                if not any(value is not None and str(value) != "" for value in values):
                    empty_scans += 1
                    if empty_scans > EXCEL_EMPTY_ROW_SCAN_LIMIT:
                        return None
                    continue
                headers = ["" if value is None else str(value) for value in values]
                while headers and not headers[-1].strip():
                    headers.pop()
                _validate_headers(headers)
                raw_columns = [[] for _ in headers]
                continue
            values = values[: len(headers)]
            if not any(value is not None for value in values):
                continue
            completed = list(values) + [None] * (len(headers) - len(values))
            rows.append(completed)
            for index, value in enumerate(completed):
                raw_columns[index].append(value)
            if len(rows) >= row_limit:
                break
        if headers is None:
            return None
        frame = _frame_from_rows(headers, rows)
        return frame, {
            name: _native_column_type(raw_columns[index])
            for index, name in enumerate(headers)
        }

    def read(
        self, path: Path, options: ReaderOptions, *, limit: int | None = None
    ) -> DatasetReadResult:
        self._validate_package(path)
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
        except (
            InvalidFileException,
            KeyError,
            OSError,
            ParseError,
            ValueError,
            zipfile.BadZipFile,
        ) as exc:
            raise ProcessingError("No fue posible leer el libro Excel XLSX.") from exc
        try:
            sheets = list(workbook.sheetnames)
            if not sheets:
                raise ProcessingError("El libro Excel no contiene hojas.")
            if options.sheet_name and options.sheet_name not in sheets:
                raise ProcessingError(
                    f"La hoja '{options.sheet_name}' no existe. "
                    f"Hojas disponibles: {', '.join(sheets)}."
                )
            selected = options.sheet_name
            parsed = None
            for candidate in [selected] if selected else sheets:
                parsed = self._sheet_frame(workbook[candidate], limit=limit)
                if parsed is not None:
                    selected = candidate
                    break
            if selected is None or parsed is None:
                raise ProcessingError("El libro Excel no contiene una hoja tabular con encabezados.")
            frame, native_schema = parsed
            _check_frame(frame, enforce_row_limit=limit is None)
            return DatasetReadResult(
                frame=frame,
                source_format=self.format,
                format_label=self.format_label,
                media_type=self.media_type,
                row_numbering="RECORD_NUMBER",
                native_schema=native_schema,
                sheets=sheets,
                selected_sheet=selected,
                row_count=frame.height if limit is None else None,
            )
        except ProcessingError:
            raise
        except (KeyError, OSError, ParseError, TypeError, ValueError) as exc:
            raise ProcessingError("El libro Excel XLSX contiene una estructura inválida.") from exc
        finally:
            workbook.close()


def _reject_json_constant(value: str) -> None:
    raise ValueError(f"Constante JSON no válida: {value}")


def _unique_json_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Clave JSON duplicada: {key}")
        result[key] = value
    return result


def _json_payload(path: Path) -> tuple[Any, str | None]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ProcessingError("El archivo JSON debe usar codificación UTF-8.") from exc
    if "\x00" in text:
        raise ProcessingError("El archivo JSON contiene bytes nulos no válidos.")
    try:
        return (
            json.loads(
                text,
                parse_float=Decimal,
                parse_constant=_reject_json_constant,
                object_pairs_hook=_unique_json_object,
            ),
            None,
        )
    except (json.JSONDecodeError, ValueError):
        # JSON Lines is a useful tabular JSON variant and remains deterministic.
        try:
            records = [
                json.loads(
                    line,
                    parse_float=Decimal,
                    parse_constant=_reject_json_constant,
                    object_pairs_hook=_unique_json_object,
                )
                for line in text.splitlines()
                if line.strip()
            ]
        except (json.JSONDecodeError, ValueError) as exc:
            raise ProcessingError("El archivo no contiene JSON válido.") from exc
        if not records:
            raise ProcessingError("El archivo JSON está vacío.")
        return records, "JSON_LINES"


def _tabular_json(payload: Any) -> tuple[list[dict[str, Any]], str | None]:
    root_key = None
    if isinstance(payload, dict):
        arrays = [(key, value) for key, value in payload.items() if isinstance(value, list)]
        if len(arrays) == 1 and all(isinstance(item, dict) for item in arrays[0][1]):
            root_key, payload = arrays[0]
        elif not arrays:
            payload = [payload]
        else:
            raise ProcessingError(
                "El JSON debe ser un arreglo de objetos o contener un único arreglo tabular."
            )
    if not isinstance(payload, list) or not payload or not all(
        isinstance(item, dict) for item in payload
    ):
        raise ProcessingError("El JSON debe contener registros tabulares representados por objetos.")
    return payload, root_key


def _flatten_json(value: Mapping[str, Any], prefix: str = "", depth: int = 0) -> dict[str, Any]:
    if depth > 8:
        raise ProcessingError("El anidamiento del JSON supera el límite soportado.")
    result: dict[str, Any] = {}
    for key, item in value.items():
        name = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(item, dict):
            nested = _flatten_json(item, name, depth + 1)
            duplicates = set(result) & set(nested)
            if duplicates:
                raise ProcessingError(
                    f"El aplanado JSON genera columnas duplicadas: {', '.join(sorted(duplicates))}."
                )
            result.update(nested)
        else:
            if name in result:
                raise ProcessingError(f"El JSON genera una columna duplicada: {name}.")
            result[name] = item
    return result


class JsonDatasetReader(DatasetReader):
    format = "JSON"
    format_label = "JSON tabular"
    media_type = "application/json"
    extensions = frozenset({".json", ".jsonl", ".ndjson"})

    def read(
        self, path: Path, options: ReaderOptions, *, limit: int | None = None
    ) -> DatasetReadResult:
        del options
        payload, variant = _json_payload(path)
        records, root_key = _tabular_json(payload)
        if limit is None and len(records) > MAX_ROWS:
            raise ProcessingError(f"Este prototipo admite hasta {MAX_ROWS:,} filas por archivo.")
        selected = records if limit is None else records[:limit]
        flattened = [_flatten_json(record) for record in selected]
        flattened_keys = {key for record in flattened for key in record}
        nested_parents = {
            key
            for key in flattened_keys
            if any(candidate.startswith(f"{key}.") for candidate in flattened_keys)
        }
        for record in flattened:
            for parent in nested_parents & record.keys():
                if record[parent] is None:
                    del record[parent]
                else:
                    raise ProcessingError(
                        f"El campo JSON '{parent}' alterna entre un objeto y un valor escalar."
                    )
        headers = list(dict.fromkeys(key for record in flattened for key in record))
        rows = [[record.get(name) for name in headers] for record in flattened]
        frame = _frame_from_rows(headers, rows)
        _check_frame(frame, enforce_row_limit=limit is None)
        result = DatasetReadResult(
            frame=frame,
            source_format=self.format,
            format_label=self.format_label,
            media_type=self.media_type,
            row_numbering="RECORD_NUMBER",
            native_schema={
                name: _native_column_type([record.get(name) for record in flattened])
                for name in frame.columns
            },
            row_count=len(records),
            metadata={
                **({"root_key": root_key} if root_key else {}),
                **({"variant": variant} if variant else {}),
            },
        )
        return result


def _parquet_metadata(path: Path) -> tuple[int, int]:
    """Read bounded footer metadata before Polars allocates decoded columns."""
    connection = duckdb.connect(":memory:")
    try:
        file_row = connection.execute(
            "SELECT num_rows FROM parquet_file_metadata(?)", [str(path)]
        ).fetchone()
        size_row = connection.execute(
            "SELECT COALESCE(SUM(total_uncompressed_size), 0) FROM parquet_metadata(?)",
            [str(path)],
        ).fetchone()
    except duckdb.Error as exc:
        raise ProcessingError("No fue posible leer la metadata del archivo Parquet.") from exc
    finally:
        connection.close()
    if not file_row or file_row[0] is None or not size_row or size_row[0] is None:
        raise ProcessingError("El archivo Parquet no contiene metadata válida.")
    row_count, expanded_bytes = int(file_row[0]), int(size_row[0])
    if row_count < 0 or expanded_bytes < 0:
        raise ProcessingError("El archivo Parquet contiene límites inválidos.")
    if row_count > MAX_ROWS:
        raise ProcessingError(f"Este prototipo admite hasta {MAX_ROWS:,} filas por archivo.")
    if expanded_bytes > MAX_PARQUET_EXPANDED_BYTES:
        raise ProcessingError(
            "El contenido expandido del archivo Parquet supera el límite de seguridad."
        )
    return row_count, expanded_bytes


class ParquetDatasetReader(DatasetReader):
    format = "PARQUET"
    format_label = "Apache Parquet"
    media_type = "application/vnd.apache.parquet"
    extensions = frozenset({".parquet", ".pq"})

    def read(
        self, path: Path, options: ReaderOptions, *, limit: int | None = None
    ) -> DatasetReadResult:
        del options
        try:
            physical_schema = pl.scan_parquet(path).collect_schema()
            _validate_headers(list(physical_schema))
            row_count, expanded_bytes = _parquet_metadata(path)
            frame = pl.read_parquet(path, n_rows=(MAX_ROWS + 1 if limit is None else limit))
        except (OSError, pl.exceptions.PolarsError) as exc:
            raise ProcessingError("No fue posible leer el archivo Parquet.") from exc
        native_schema = {name: str(data_type) for name, data_type in physical_schema.items()}
        frame = _string_frame(frame)
        _check_frame(frame, enforce_row_limit=limit is None)
        return DatasetReadResult(
            frame=frame,
            source_format=self.format,
            format_label=self.format_label,
            media_type=self.media_type,
            row_numbering="RECORD_NUMBER",
            native_schema=native_schema,
            row_count=row_count,
            metadata={"expanded_size_bytes": expanded_bytes},
        )


class DatasetReaderRegistry:
    """Resolve readers by strong file signatures first, then by source name."""

    def __init__(self, readers: list[DatasetReader] | None = None):
        self._readers: dict[str, DatasetReader] = {}
        self._extensions: dict[str, DatasetReader] = {}
        for reader in readers or []:
            self.register(reader)

    def register(self, reader: DatasetReader) -> None:
        if reader.format in self._readers:
            raise ValueError(f"Ya existe un lector para {reader.format}.")
        self._readers[reader.format] = reader
        for extension in reader.extensions:
            if extension in self._extensions:
                raise ValueError(f"Ya existe un lector para {extension}.")
            self._extensions[extension] = reader

    @property
    def formats(self) -> list[dict[str, Any]]:
        return [
            {
                "format": reader.format,
                "label": reader.format_label,
                "media_type": reader.media_type,
                "extensions": sorted(reader.extensions),
            }
            for reader in self._readers.values()
        ]

    def resolve(self, path: Path, filename: str | None = None) -> DatasetReader:
        try:
            with path.open("rb") as source:
                prefix = source.read(4096)
                source.seek(0, 2)
                if source.tell() >= 4:
                    source.seek(-4, 2)
                    suffix = source.read(4)
                else:
                    suffix = b""
        except (OSError, ValueError) as exc:
            raise ProcessingError("No fue posible inspeccionar el archivo cargado.") from exc
        if prefix[:4] == b"PAR1" and suffix == b"PAR1":
            return self._readers["PARQUET"]
        if prefix[:2] == b"PK":
            return self._readers["XLSX"]
        try:
            sample = prefix.decode("utf-8-sig").lstrip()
        except UnicodeDecodeError:
            sample = ""
        if sample.startswith(("{", "[")):
            return self._readers["JSON"]
        extension = Path(filename or path.name).suffix.lower()
        if extension in self._extensions:
            return self._extensions[extension]
        if sample:
            try:
                delimiter = _detect_delimiter(sample, csv_legacy_default=False)
                if len(_headers_from_text(sample, delimiter)) > 1:
                    return self._readers["TXT"]
            except ProcessingError:
                pass
        raise UnsupportedDatasetFormat(
            "Formato no soportado. Usa CSV, Excel XLSX, JSON, Parquet o TXT delimitado."
        )

    def read(
        self,
        path: Path,
        filename: str | None = None,
        options: Mapping[str, Any] | ReaderOptions | None = None,
        *,
        inspect: bool = False,
    ) -> DatasetReadResult:
        parsed_options = (
            options if isinstance(options, ReaderOptions) else ReaderOptions.from_mapping(options)
        )
        reader = self.resolve(path, filename)
        return reader.inspect(path, parsed_options) if inspect else reader.read(path, parsed_options)


dataset_reader_registry = DatasetReaderRegistry(
    [
        CsvDatasetReader(),
        ExcelDatasetReader(),
        JsonDatasetReader(),
        ParquetDatasetReader(),
        TxtDatasetReader(),
    ]
)


@dataclass(frozen=True)
class LocalFileDatasetSource:
    """Local-file source adapter backed by the registered format readers."""

    path: Path
    filename: str | None = None
    registry: DatasetReaderRegistry | None = None
    source_kind: str = field(default="LOCAL_FILE", init=False)

    def read(
        self,
        options: ReaderOptions | Mapping[str, Any] | None = None,
        *,
        inspect: bool = False,
    ) -> DatasetReadResult:
        registry = self.registry or dataset_reader_registry
        return registry.read(self.path, self.filename, options, inspect=inspect)


def read_dataset(
    source: Path | DatasetSource,
    filename: str | None = None,
    options: Mapping[str, Any] | ReaderOptions | None = None,
) -> DatasetReadResult:
    adapter = LocalFileDatasetSource(source, filename) if isinstance(source, Path) else source
    return adapter.read(options)


def inspect_dataset(
    source: Path | DatasetSource,
    filename: str | None = None,
    options: Mapping[str, Any] | ReaderOptions | None = None,
) -> DatasetReadResult:
    adapter = LocalFileDatasetSource(source, filename) if isinstance(source, Path) else source
    return adapter.read(options, inspect=True)


def delimited_record_lines(path: Path, delimiter: str) -> list[int]:
    """Return physical start lines for CSV/TXT records, including multiline records."""
    with path.open(encoding="utf-8-sig", newline="") as source:
        records = csv.reader(source, delimiter=delimiter)
        next(records, None)
        lines: list[int] = []
        previous_line = records.line_num
        for _ in records:
            lines.append(previous_line + 1)
            previous_line = records.line_num
    return lines
