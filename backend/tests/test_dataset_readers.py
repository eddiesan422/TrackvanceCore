import io
import json
import zipfile
from decimal import Decimal

import polars as pl
import pytest
from openpyxl import Workbook
from sqlalchemy import select

from trackvance import dataset_readers as reader_module
from trackvance.dataset_readers import (
    ProcessingError,
    ReaderOptions,
    dataset_reader_registry,
    inspect_dataset,
    read_dataset,
)
from trackvance.models import Artifact, AuditEvent, DatasetVersion
from trackvance.worker import process_once


def workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Resumen"
    first.append(["document_id", "amount", "transaction_date"])
    first.append(["001234", Decimal("9007199254740993.01"), "2026-09-15"])
    second = workbook.create_sheet("Detalle")
    second.append(["code", "quantity", "calculation"])
    second.append(["A-01", 2, "=1+1"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_registry_exposes_extensible_reader_catalog():
    assert [item["format"] for item in dataset_reader_registry.formats] == [
        "CSV",
        "XLSX",
        "JSON",
        "PARQUET",
        "TXT",
    ]
    assert dataset_reader_registry.formats[-1]["extensions"] == [".tsv", ".txt"]


def test_txt_detects_or_accepts_delimiter_and_preserves_observed_values(tmp_path):
    path = tmp_path / "customers.txt"
    path.write_text("customer_id|name|amount\n001|  Cliente 12  |10.50\n", encoding="utf-8")

    detected = read_dataset(path, path.name)
    explicit = read_dataset(path, path.name, {"delimiter": "|"})

    assert detected.source_format == "TXT"
    assert detected.detected_delimiter == "|"
    assert detected.row_numbering == "PHYSICAL_LINE"
    assert detected.frame.to_dicts() == [
        {"customer_id": "001", "name": "  Cliente 12  ", "amount": "10.50"}
    ]
    assert explicit.frame.equals(detected.frame)


def test_txt_requires_delimiter_when_it_cannot_be_detected(tmp_path):
    path = tmp_path / "single-column.txt"
    path.write_text("value\nA\n", encoding="utf-8")

    with pytest.raises(ProcessingError, match="delimitador"):
        read_dataset(path, path.name)

    result = read_dataset(path, path.name, {"delimiter": ";"})
    assert result.frame.to_dicts() == [{"value": "A"}]


def test_excel_inspection_lists_sheets_and_selected_sheet_is_deterministic(tmp_path):
    path = tmp_path / "orders.xlsx"
    path.write_bytes(workbook_bytes())

    inspection = inspect_dataset(path, path.name)
    selected = read_dataset(path, path.name, {"sheet_name": "Detalle"})

    assert inspection.sheets == ["Resumen", "Detalle"]
    assert inspection.selected_sheet == "Resumen"
    assert inspection.frame["document_id"].to_list() == ["001234"]
    assert selected.selected_sheet == "Detalle"
    assert selected.frame.to_dicts() == [
        {"code": "A-01", "quantity": "2", "calculation": "=1+1"}
    ]


def test_excel_rejects_unknown_sheet_with_available_names(tmp_path):
    path = tmp_path / "orders.xlsx"
    path.write_bytes(workbook_bytes())

    with pytest.raises(ProcessingError, match="Hojas disponibles: Resumen, Detalle"):
        read_dataset(path, path.name, {"sheet_name": "Missing"})


def test_excel_rejects_a_malformed_workbook_as_invalid_data(tmp_path):
    path = tmp_path / "malformed.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "invalid")
        archive.writestr("xl/workbook.xml", "invalid")

    with pytest.raises(ProcessingError, match="Excel XLSX"):
        read_dataset(path, path.name)


def test_json_flattens_simple_nesting_without_float_roundtrip(tmp_path):
    path = tmp_path / "orders.json"
    path.write_text(
        '{"records":['
        '{"id":"001","amount":9007199254740993.01,'
        '"customer":{"name":"Ángela","address":{"city":"Bogotá"}}},'
        '{"id":"002","amount":null,"customer":{"name":"  Cliente 12  "}}'
        '] ,"source":"local"}',
        encoding="utf-8",
    )

    result = read_dataset(path, path.name)

    assert result.frame.columns == [
        "id",
        "amount",
        "customer.name",
        "customer.address.city",
    ]
    assert result.frame.to_dicts() == [
        {
            "id": "001",
            "amount": "9007199254740993.01",
            "customer.name": "Ángela",
            "customer.address.city": "Bogotá",
        },
        {
            "id": "002",
            "amount": None,
            "customer.name": "  Cliente 12  ",
            "customer.address.city": None,
        },
    ]
    assert result.metadata == {"root_key": "records"}
    assert result.native_schema["amount"] == "Decimal"


def test_json_treats_a_null_nested_object_as_null_children(tmp_path):
    path = tmp_path / "customers.json"
    path.write_text(
        '[{"id":"A","customer":{"name":"Ángela"}},{"id":"B","customer":null}]',
        encoding="utf-8",
    )

    result = read_dataset(path, path.name)

    assert result.frame.columns == ["id", "customer.name"]
    assert result.frame.to_dicts()[1] == {"id": "B", "customer.name": None}


def test_parquet_uses_embedded_schema_and_normalizes_scalars(tmp_path):
    path = tmp_path / "orders.parquet"
    pl.DataFrame(
        {
            "document_id": [1234567, 7654321],
            "amount": [Decimal("10.25"), Decimal("20.50")],
        }
    ).write_parquet(path)

    result = read_dataset(path, "misleading.csv")

    assert result.source_format == "PARQUET"
    assert result.native_schema["document_id"].startswith("Int")
    assert result.native_schema["amount"].startswith("Decimal")
    assert result.frame.schema == {"document_id": pl.String, "amount": pl.String}
    assert result.frame.to_dicts()[0] == {"document_id": "1234567", "amount": "10.25"}


def test_parquet_rejects_unsafe_expansion_before_decoding_rows(tmp_path, monkeypatch):
    path = tmp_path / "compressed.parquet"
    pl.DataFrame({"payload": ["A" * 10_000] * 20}).write_parquet(path)
    monkeypatch.setattr(reader_module, "MAX_PARQUET_EXPANDED_BYTES", 1)

    def fail_decode(*_args, **_kwargs):
        raise AssertionError("Polars must not decode rows after the footer exceeds the limit")

    monkeypatch.setattr(pl, "read_parquet", fail_decode)
    with pytest.raises(ProcessingError, match="contenido expandido"):
        read_dataset(path, path.name)


def test_detected_format_controls_original_artifact_mime_and_download_suffix(
    authenticated,
):
    stream = io.BytesIO()
    pl.DataFrame({"record_key": ["A"], "amount": [1]}).write_parquet(stream)
    dataset = authenticated.post(
        "/api/v1/datasets", json={"name": "Misleading extension"}
    ).json()

    upload = authenticated.post(
        f"/api/v1/datasets/{dataset['id']}/versions/upload",
        files={"file": ("renamed.csv", stream.getvalue(), "text/csv")},
    )

    assert upload.status_code == 201, upload.text
    payload = upload.json()
    assert payload["ingestion_metadata"]["source_format"] == "PARQUET"
    original = next(
        artifact
        for artifact in payload["artifacts"]
        if artifact["artifact_id"] == payload["original_artifact_id"]
    )
    assert original["media_type"] == "application/vnd.apache.parquet"
    download = authenticated.get(
        f"/api/v1/artifacts/{payload['original_artifact_id']}/download"
    )
    assert download.status_code == 200
    assert download.headers["content-type"] == "application/vnd.apache.parquet"
    assert ".parquet" in download.headers["content-disposition"]


def test_parquet_null_only_columns_keep_embedded_types_for_rule_selectors(
    authenticated,
):
    stream = io.BytesIO()
    pl.DataFrame(
        {
            "quantity": pl.Series([None], dtype=pl.Int64),
            "amount": pl.Series([None], dtype=pl.Decimal(18, 2)),
            "booked_on": pl.Series([None], dtype=pl.Date),
            "observed_at": pl.Series([None], dtype=pl.Datetime),
        }
    ).write_parquet(stream)

    response = authenticated.post(
        "/api/v1/datasets/uploads/inspect",
        files={"file": ("typed.parquet", stream.getvalue(), "application/vnd.apache.parquet")},
    )

    assert response.status_code == 200, response.text
    columns = {column["name"]: column for column in response.json()["columns"]}
    assert columns["quantity"]["logical_type"] == "INT64"
    assert columns["quantity"]["numeric"] is True
    assert columns["amount"]["logical_type"] == "DECIMAL"
    assert columns["amount"]["numeric"] is True
    assert columns["booked_on"]["logical_type"] == "DATE"
    assert columns["observed_at"]["logical_type"] == "TIMESTAMP"


@pytest.mark.parametrize(
    ("filename", "content_type", "content", "options", "expected_format", "expected_columns"),
    [
        ("data.csv", "text/csv", b"id,amount\n001,10.5\n", {}, "CSV", ["id", "amount"]),
        (
            "data.txt",
            "text/plain",
            b"id\tamount\n001\t10.5\n",
            {"delimiter": "\\t"},
            "TXT",
            ["id", "amount"],
        ),
        (
            "data.json",
            "application/json",
            b'[{"id":"001","amount":10.5}]',
            {},
            "JSON",
            ["id", "amount"],
        ),
    ],
)
def test_inspection_endpoint_detects_text_formats(
    authenticated,
    filename,
    content_type,
    content,
    options,
    expected_format,
    expected_columns,
):
    response = authenticated.post(
        "/api/v1/datasets/uploads/inspect",
        files={"file": (filename, content, content_type)},
        data={"reader_options": json.dumps(options)},
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["format"] == expected_format
    assert [item["name"] for item in payload["columns"]] == expected_columns
    assert payload["sampled_rows"] == 1
    assert len(payload["supported_formats"]) == 5


def test_inspection_endpoint_lists_excel_sheets(authenticated):
    response = authenticated.post(
        "/api/v1/datasets/uploads/inspect",
        files={
            "file": (
                "orders.xlsx",
                workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"reader_options": json.dumps({"sheet_name": "Detalle"})},
    )

    assert response.status_code == 200, response.text
    assert response.json()["sheets"] == ["Resumen", "Detalle"]
    assert response.json()["selected_sheet"] == "Detalle"
    assert response.json()["reader_options"] == {"sheet_name": "Detalle"}


def test_upload_persists_reader_metadata_artifacts_profile_and_audit(
    authenticated, database
):
    dataset = authenticated.post("/api/v1/datasets", json={"name": "Excel orders"}).json()
    response = authenticated.post(
        f"/api/v1/datasets/{dataset['id']}/versions/upload",
        files={
            "file": (
                "orders.xlsx",
                workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"reader_options": json.dumps({"sheet_name": "Detalle"})},
    )

    assert response.status_code == 201, response.text
    payload = response.json()
    metadata = payload["ingestion_metadata"]
    assert metadata["reader"] == {"key": "XLSX", "version": 1}
    assert metadata["source_format"] == "XLSX"
    assert metadata["reader_options"] == {"sheet_name": "Detalle"}
    assert metadata["row_numbering"] == "RECORD_NUMBER"
    assert payload["profile"]["row_numbering"] == "RECORD_NUMBER"
    assert payload["schema"][1]["logical_type"] == "INT64"
    sample = authenticated.get(f"/api/v1/dataset-versions/{payload['id']}/profile")
    assert sample.status_code == 200
    assert sample.json()["sample"] == [
        {"code": "A-01", "quantity": "2", "calculation": "=1+1"}
    ]
    with database() as db:
        version = db.get(DatasetVersion, payload["id"])
        original = db.get(Artifact, version.original_artifact_id)
        event = db.scalar(
            select(AuditEvent).where(
                AuditEvent.event_type == "DATASET_UPLOADED",
                AuditEvent.subject_id == version.id,
            )
        )
        assert original.media_type == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert event.metadata_json["format"] == "XLSX"
        assert event.metadata_json["sheet_name"] == "Detalle"


def test_upload_parquet_and_txt_keep_correct_record_numbering(authenticated):
    parquet_stream = io.BytesIO()
    pl.DataFrame({"id": ["A", "B"], "amount": [1, 2]}).write_parquet(parquet_stream)
    for name, content, options, expected_format, numbering in [
        ("source.parquet", parquet_stream.getvalue(), {}, "PARQUET", "RECORD_NUMBER"),
        ("source.txt", b"id|amount\nA|1\nB|2\n", {}, "TXT", "PHYSICAL_LINE"),
    ]:
        dataset = authenticated.post(
            "/api/v1/datasets", json={"name": f"Dataset {expected_format}"}
        ).json()
        response = authenticated.post(
            f"/api/v1/datasets/{dataset['id']}/versions/upload",
            files={"file": (name, content, "application/octet-stream")},
            data={"reader_options": json.dumps(options)},
        )
        assert response.status_code == 201, response.text
        assert response.json()["ingestion_metadata"]["source_format"] == expected_format
        assert response.json()["profile"]["row_numbering"] == numbering


def test_unknown_binary_format_has_specific_error(authenticated):
    response = authenticated.post(
        "/api/v1/datasets/uploads/inspect",
        files={"file": ("data.bin", b"\x00\x01\x02\x03\x04", "application/octet-stream")},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "UNSUPPORTED_FORMAT"


def test_delimited_text_without_a_known_extension_is_detected_from_content(authenticated):
    response = authenticated.post(
        "/api/v1/datasets/uploads/inspect",
        files={
            "file": (
                "daily-feed.dat",
                b"record_key|amount\nA|10.5\n",
                "application/octet-stream",
            )
        },
    )

    assert response.status_code == 200, response.text
    assert response.json()["format"] == "TXT"
    assert response.json()["detected_delimiter"] == "|"


def test_reader_options_reject_unknown_or_unsafe_values():
    with pytest.raises(ProcessingError, match="no soportadas"):
        ReaderOptions.from_mapping({"password": "secret"})
    with pytest.raises(ProcessingError, match="ASCII"):
        ReaderOptions.from_mapping({"delimiter": "ab"})


def test_all_file_readers_feed_the_same_intake_engine_contract(authenticated):
    rows = [
        {"record_key": "A", "amount": "10"},
        {"record_key": "B", "amount": "-2"},
        {"record_key": "C", "amount": "bad"},
    ]
    parquet = io.BytesIO()
    pl.DataFrame(rows).write_parquet(parquet)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["record_key", "amount"])
    for row in rows:
        worksheet.append([row["record_key"], row["amount"]])
    excel = io.BytesIO()
    workbook.save(excel)
    workbook.close()
    sources = [
        ("source.csv", b"record_key,amount\nA,10\nB,-2\nC,bad\n", "text/csv", {}),
        (
            "source.txt",
            b"record_key|amount\nA|10\nB|-2\nC|bad\n",
            "text/plain",
            {"delimiter": "|"},
        ),
        (
            "source.json",
            json.dumps(rows).encode(),
            "application/json",
            {},
        ),
        (
            "source.parquet",
            parquet.getvalue(),
            "application/vnd.apache.parquet",
            {},
        ),
        (
            "source.xlsx",
            excel.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            {},
        ),
    ]
    observed = []
    for index, (filename, content, content_type, options) in enumerate(sources):
        dataset = authenticated.post(
            "/api/v1/datasets", json={"name": f"Reader parity {index}"}
        ).json()
        upload = authenticated.post(
            f"/api/v1/datasets/{dataset['id']}/versions/upload",
            files={"file": (filename, content, content_type)},
            data={"reader_options": json.dumps(options)},
        )
        assert upload.status_code == 201, upload.text
        contract = authenticated.post(
            "/api/v1/intake/contracts",
            json={
                "name": f"Reader parity {index}",
                "dataset_id": dataset["id"],
                "config": {
                    "numeric_columns": ["amount"],
                    "positive_columns": ["amount"],
                    "max_error_rate": 1,
                },
            },
        )
        assert contract.status_code == 201, contract.text
        queued = authenticated.post(
            "/api/v1/intake/runs",
            json={
                "contract_id": contract.json()["id"],
                "dataset_version_id": upload.json()["id"],
            },
        )
        assert queued.status_code == 202, queued.text
        assert process_once(f"reader-parity-{index}") is True
        completed = authenticated.get(f"/api/v1/runs/{queued.json()['id']}")
        assert completed.status_code == 200
        metrics = completed.json()["metrics"]
        observed.append(
            {
                "total_rows": metrics["total_rows"],
                "valid_rows": metrics["valid_rows"],
                "error_rows": metrics["error_rows"],
                "rules": [
                    (rule["code"], rule["failed_count"])
                    for rule in metrics["rules"]
                ],
            }
        )

    assert observed == [observed[0]] * len(sources)
    assert observed[0] == {
        "total_rows": 3,
        "valid_rows": 1,
        "error_rows": 2,
        "rules": [("NUMERIC", 1), ("POSITIVE", 2)],
    }
