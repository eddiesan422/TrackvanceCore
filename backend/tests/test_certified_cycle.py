import io
from collections import Counter

from certified_data import COLUMNS, INTAKE_CONFIG, MONITOR_CONFIG, csv_bytes, original_rows
from openpyxl import load_workbook
from sqlalchemy import select

from trackvance.models import AuditEvent, User
from trackvance.worker import process_once


def upload(client, dataset, data, overrides=None):
    response = client.post(f"/api/v1/datasets/{dataset}/versions/upload", files={"file": ("certified.csv", data, "text/csv")}, data=overrides or {})
    assert response.status_code == 201, response.text
    return response.json()


def execute(client, endpoint, payload, key=None):
    queued = client.post('/api/v1' + endpoint, json=payload, headers={"Idempotency-Key": key} if key else {})
    assert queued.status_code == 202, queued.text
    run_id = queued.json()["id"]
    assert process_once("certification-worker")
    run = client.get(f"/api/v1/runs/{run_id}").json()
    assert run["status"] == "SUCCESS", run
    return run


def test_certified_business_cycle_and_reports(authenticated, database):
    client = authenticated
    rows = original_rows()
    dataset = client.post('/api/v1/datasets', json={"name": "Certificación original"}).json()["id"]
    version = upload(client, dataset, csv_bytes(rows))
    columns = {c["name"]: c for c in version["profile"]["columns"]}
    assert columns["customer_name"]["distinct_count"] == 44
    assert columns["transaction_date"]["logical_type"] == "DATE"
    assert columns["document_id"]["logical_type"] == "STRING"
    config = client.post('/api/v1/intake/contracts', json={"name": "IT01", "dataset_id": dataset, "config": INTAKE_CONFIG})
    assert config.status_code == 201, config.text
    config = config.json()
    intake = execute(client, '/intake/runs', {"contract_id": config["id"], "dataset_version_id": version["id"]}, "certified-intake")
    metrics = intake["metrics"]
    assert (metrics["total_rows"], metrics["valid_rows"], metrics["error_rows"], metrics["acceptance_rate"], intake["decision"]) == (120, 110, 10, 91.67, "REJECTED")
    counts = Counter()
    for rule in metrics["rules"]:
        counts[rule["code"]] += rule["failed_count"]
    assert counts["REQUIRED"] == 4 and counts["UNIQUE"] == 4 and counts["POSITIVE"] == 2
    output = client.get(f"/api/v1/dataset-versions/{intake['output_version_id']}/profile").json()
    assert output["source_type"] == "INTAKE_OUTPUT" and output["original_artifact_id"] is None
    assert output["canonical_artifact_id"] and output["source_run_id"] == intake["id"]
    replay = client.post('/api/v1/intake/runs', json={"contract_id": config["id"], "dataset_version_id": version["id"]}, headers={"Idempotency-Key": "certified-intake"})
    assert replay.json()["id"] == intake["id"]
    recon_config = client.post('/api/v1/recon/controls', json={"name": "Conciliación certificada", "dataset_id": dataset, "target_dataset_id": output["dataset_id"], "config": {"key_columns": ["transaction_id"], "amount_column": "total_amount", "tolerance": "0"}})
    assert recon_config.status_code == 201, recon_config.text
    recon = execute(client, '/recon/runs', {"control_id": recon_config.json()["id"], "source_version_id": version["id"], "target_version_id": output["id"]})
    result = client.get(f"/api/v1/runs/{recon['id']}/results?limit=1000").json()
    assert Counter(r["classification"] for r in result["items"]) == {"MATCH": 110, "SOURCE_ONLY": 6, "DUPLICATE_SOURCE": 4}
    assert recon["metrics"]["match_rate"] == 91.67
    monitor = client.post('/api/v1/monitors', json={"name": "Monitor certificado", "dataset_id": output["dataset_id"], "config": MONITOR_CONFIG})
    assert monitor.status_code == 201, monitor.text
    monitor_id = monitor.json()["id"]
    baseline = execute(client, f'/monitors/{monitor_id}/runs', {"dataset_version_id": output["id"]})
    assert (baseline["metrics"]["total_checks"], baseline["metrics"]["failed_checks"], baseline["metrics"]["health_score"]) == (9, 0, 100)
    v2_rows = [dict(r) for r in rows[:90]]
    for i, col in enumerate(["customer_id", "email", "city"]):
        v2_rows[i][col] = None
    v2 = upload(client, output["dataset_id"], csv_bytes(v2_rows))
    sentinel2 = execute(client, f'/monitors/{monitor_id}/runs', {"dataset_version_id": v2["id"]})
    assert (sentinel2["metrics"]["failed_checks"], sentinel2["metrics"]["health_score"]) == (4, 55.56)
    v3_rows = [{k: v for k, v in r.items() if k != "source_system"} for r in rows[:90]]
    v3 = upload(client, output["dataset_id"], csv_bytes(v3_rows, [c for c in COLUMNS if c != "source_system"]))
    sentinel3 = execute(client, f'/monitors/{monitor_id}/runs', {"dataset_version_id": v3["id"]})
    assert (sentinel3["metrics"]["failed_checks"], sentinel3["metrics"]["health_score"]) == (1, 88.89)
    for run, sheets in [(intake, ["Resumen", "Errores", "Reglas", "Trazabilidad"]), (recon, ["Resumen", "Resultados", "Hallazgos", "Trazabilidad"]), (sentinel3, ["Resumen", "Controles", "Hallazgos", "Trazabilidad"])]:
        response = client.get(f"/api/v1/runs/{run['id']}/export.xlsx")
        assert response.status_code == 200, response.text
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert f"trackvance_{run['module']}_{run['id']}.xlsx" in response.headers["content-disposition"]
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == sheets
        assert all(c.data_type != "f" for sheet in workbook for row in sheet for c in row)
    finding_id = recon["findings"][0]["id"]
    case = client.post(f'/api/v1/findings/{finding_id}/exceptions').json()
    case = client.patch(f"/api/v1/exceptions/{case['id']}", json={"version": case["version"], "state": "INVESTIGATING", "owner": "Test User"}).json()
    case = client.patch(f"/api/v1/exceptions/{case['id']}", json={"version": case["version"], "state": "RESOLVED", "root_cause": "Duplicados en la fuente", "resolution": "Corregir próxima entrega"}).json()
    assert case["version"] == 3 and case["state"] == "RESOLVED" and case["run_id"] == recon["id"] and len(case["events"]) == 3
    with database() as db:
        events = db.scalars(select(AuditEvent)).all()
        kinds = {e.event_type for e in events}
        assert {"DATASET_UPLOADED", "CONFIGURATION_PUBLISHED", "RUN_QUEUED", "RUN_COMPLETED", "EXCEPTION_CREATED", "EXCEPTION_UPDATED", "EXPORT_DOWNLOADED"} <= kinds
        exports = [e for e in events if e.event_type == "EXPORT_DOWNLOADED"]
        assert all(e.actor_type == "USER" and e.actor_id == "test-user" and e.request_id and e.run_id for e in exports)


def test_role_permissions_are_enforced_server_side(authenticated, database):
    client = authenticated
    for role, expected in [("Auditor", 403), ("Operations", 403), ("Data Analyst", 201), ("Data Owner", 201), ("Administrator", 201), ("unrecognized", 403)]:
        with database() as db:
            db.get(User, "test-user").role = role
            db.commit()
        assert client.post('/api/v1/datasets', json={"name": role}).status_code == expected
