"""Semantic, package-budget and publication tests for incremental XLSX."""

import hashlib
import importlib.util
import json
import re
import struct
import zipfile
from dataclasses import replace
from datetime import date, datetime, timedelta
from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from trackvance import acquisition
from trackvance.acquisition_config import AcquisitionLimits
from trackvance.acquisition_errors import AcquisitionReadError, processing_diagnostic
from trackvance.acquisition_models import AcquisitionRun, AcquisitionUpload
from trackvance.artifactstore import file_hash
from trackvance.automation_models import OutboxEvent
from trackvance.batch_readers import FileBatchReader, inspect_file
from trackvance.dataset_readers import ExcelDatasetReader, ReaderOptions
from trackvance.dataset_scans import iter_version_batches, profile_paths
from trackvance.db import utcnow
from trackvance.events import consume_notification
from trackvance.models import DatasetVersion, Job
from trackvance.processing import ProcessingError
from trackvance.xlsx_streaming import XlsxStreamingReader, _Index

_fixture_spec = importlib.util.spec_from_file_location("xlsx_independent_fixture", Path(__file__).resolve().parents[2] / "scripts" / "xlsx_fixtures.py")
assert _fixture_spec is not None and _fixture_spec.loader is not None
fixtures = importlib.util.module_from_spec(_fixture_spec)
_fixture_spec.loader.exec_module(fixtures)


def modify_package(source, target, transform):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, "x", compression=zipfile.ZIP_DEFLATED) as changed:
        for item in original.infolist():
            changed.writestr(item.filename, transform(item.filename, original.read(item.filename)))
    return target


@pytest.mark.parametrize("strings", ["inline", "shared"])
def test_selected_sheet_wrong_dimension_blanks_dates_formulas_unicode_and_exact_decimal(tmp_path, strings):
    fixture = fixtures.generate(tmp_path, 137, strings=strings)
    path = Path(fixture["path"])
    digest = hashlib.sha256()
    reader = FileBatchReader(path, path.name, {"sheet_name": "Datos"}, replace(AcquisitionLimits(), batch_rows=7))
    count, paths = 0, []
    for index, batch in enumerate(reader):
        assert batch.frame.height <= 7
        assert batch.record_numbers == list(range(count + 5, count + batch.frame.height + 5))
        assert batch.row_numbering == "PHYSICAL_SHEET_ROW"
        for row in batch.frame.iter_rows(named=True):
            count += 1
            assert row == fixtures.expected_row(count)
            digest.update(fixtures.canonical(row))
        part = tmp_path / f"part-{index}.parquet"
        batch.physical_frame().write_parquet(part)
        paths.append(part)
    assert count == 137 and reader.total_rows == count
    assert digest.hexdigest() == fixture["canonical_rows_sha256"]
    assert file_hash(path) == fixture["file_sha256"]
    schema, profile, _ = profile_paths(paths, native_types=reader.native_schema)
    assert profile["row_count"] == 137
    assert next(column for column in schema if column["name"] == "amount")["logical_type"] == "DECIMAL"
    assert next(column for column in schema if column["name"] == "date")["logical_type"] == "DATE"
    inspection = inspect_file(path, path.name, {"sheet_name": "Datos"})
    assert inspection["sampled_rows"] == 100 and inspection["row_count"] is None
    assert inspection["effective_limits"]["header_row_number"] == 4
    assert inspection["sheets"] == ["Información", "Datos"]
    assert inspection["sample"][1]["observed"] == fixtures.expected_row(2)["observed"]


def test_100001_rows_complete_inference_and_hash_independent_of_legacy_bounds(tmp_path, monkeypatch):
    monkeypatch.setenv("TRACKVANCE_ACQUISITION_BOUNDED_FORMAT_BYTES", "10485760")
    monkeypatch.setenv("TRACKVANCE_ACQUISITION_BOUNDED_FORMAT_ROWS", "100000")
    fixture = fixtures.generate(tmp_path, 100001, strings="inline", wrong_dimension=False)
    path = Path(fixture["path"])
    reader = FileBatchReader(path, path.name, {"sheet_name": "Datos"})
    digest, rows, batch_count = hashlib.sha256(), 0, 0
    parts = []
    for batch in reader:
        batch_count += 1
        assert batch.frame.height <= 5000
        for row in batch.frame.iter_rows(named=True):
            digest.update(fixtures.canonical(row))
            rows += 1
        part = tmp_path / f"population-{batch_count}.parquet"
        batch.physical_frame().write_parquet(part)
        parts.append(part)
    assert rows == 100001 and batch_count == 21
    assert reader.native_schema["late_type"] == "String"
    assert digest.hexdigest() == fixture["canonical_rows_sha256"]
    schema, profile, _ = profile_paths(parts, native_types=reader.native_schema)
    assert profile["row_count"] == 100001
    late_schema = next(column for column in schema if column["name"] == "late_type")
    late_profile = next(column for column in profile["columns"] if column["name"] == "late_type")
    assert late_schema["logical_type"] == "STRING" and late_profile["distinct_count"] == 998
    assert inspect_file(path, path.name, {"sheet_name": "Datos"})["columns"][-1]["logical_type"] == "INT64"
    # The normal async change does not elevate the synchronous reader's cap.
    with pytest.raises(ProcessingError, match="100,000"):
        ExcelDatasetReader().read(path, ReaderOptions(sheet_name="Datos"))


def test_shared_strings_inspection_stops_at_budget_without_full_index_or_false_empty(tmp_path, monkeypatch):
    fixture = fixtures.generate(tmp_path, 20, strings="shared")
    path = Path(fixture["path"])
    changed = tmp_path / "late-sst-header.xlsx"
    # Add many valid unused shared strings before the original SST entries and
    # shift every reference. Full worker can index it; HTTP cannot scan it all.
    prefix_count = 200
    def transform(name, content):
        if name == "xl/sharedStrings.xml":
            at = content.index(b">") + 1
            return content[:at] + (b"<si><t>" + b"x" * 1000 + b"</t></si>") * prefix_count + content[at:]
        if name.startswith("xl/worksheets/"):
            import re
            return re.sub(br'(<c[^>]* t="s"><v>)(\d+)(</v>)', lambda match: match[1] + str(int(match[2]) + prefix_count).encode() + match[3], content)
        return content
    modify_package(path, changed, transform)
    monkeypatch.setattr(_Index, "__init__", lambda *_: pytest.fail("HTTP must not create the whole SST disk index"))
    inspection = inspect_file(changed, changed.name, {"sheet_name": "Datos"}, replace(AcquisitionLimits(), xlsx_inspection_bytes=65536))
    assert inspection["inspection_limited"] is True
    assert inspection["row_count"] is None and inspection["columns"] == [] and inspection["sample"] == []
    assert inspection["effective_limits"]["inspection_limited"] is True


@pytest.mark.parametrize("change,code", [
    (lambda name, data: b'<!DOCTYPE worksheet [<!ENTITY attack "secret">]>' + data if name.endswith("sheet2.xml") else data, "ACQUISITION_XLSX_INVALID_STRUCTURE"),
    (lambda name, data: data.replace(b"<sheetData>", b'<sheetData bad="' + b"x" * 100000 + b'">') if name.endswith("sheet2.xml") else data, "ACQUISITION_XLSX_INVALID_STRUCTURE"),
])
def test_hostile_xml_is_rejected_before_entity_or_large_token_allocation(tmp_path, change, code):
    fixture = fixtures.generate(tmp_path, 3)
    changed = modify_package(Path(fixture["path"]), tmp_path / "hostile.xlsx", change)
    with pytest.raises(AcquisitionReadError) as error:
        list(FileBatchReader(changed, changed.name, {"sheet_name": "Datos"}))
    assert error.value.code == code and "secret" not in error.value.message


@pytest.mark.parametrize("field,value,code", [
    ("xlsx_max_expanded_bytes", 1024, "ACQUISITION_EXPANDED_SIZE_LIMIT"),
    ("xlsx_metadata_bytes", 10, "ACQUISITION_XLSX_METADATA_LIMIT"),
    ("xlsx_max_cells", 9, "ACQUISITION_XLSX_CELL_LIMIT"),
    ("xlsx_max_record_bytes", 10, "ACQUISITION_RECORD_LIMIT"),
    ("xlsx_max_styles", 1, "ACQUISITION_XLSX_STYLE_LIMIT"),
    ("xlsx_temp_bytes", 1, "ACQUISITION_TEMP_DISK_LIMIT"),
])
def test_distinct_xlsx_budgets_report_structured_effective_limit(tmp_path, field, value, code):
    fixture = fixtures.generate(tmp_path, 3, strings="shared")
    path = Path(fixture["path"])
    with pytest.raises(AcquisitionReadError) as error:
        list(FileBatchReader(path, path.name, {"sheet_name": "Datos"}, replace(AcquisitionLimits(), **{field: value})))
    assert error.value.code == code and error.value.details["maximum"] == value
    assert str(path) not in error.value.message


def test_incremental_indexing_checks_cancel_and_closes_resources_before_rows(tmp_path, monkeypatch):
    fixture = fixtures.generate(tmp_path, 1000, strings="shared")
    puts = [0]
    put = _Index.put
    def count_put(index, position, value):
        puts[0] += 1
        put(index, position, value)
    monkeypatch.setattr(_Index, "put", count_put)
    def checkpoint():
        if puts[0] > 200:
            raise acquisition.AcquisitionStopped("ACQUISITION_CANCELLED", "Cancelación solicitada.")
    source = XlsxStreamingReader(Path(fixture["path"]), ReaderOptions(sheet_name="Datos"), AcquisitionLimits(), checkpoint)
    with pytest.raises(acquisition.AcquisitionStopped):
        list(source)
    assert 200 < puts[0] < fixture["shared_string_count"]
    assert source.index is not None
    assert not source.index.path.exists()
    with pytest.raises(Exception, match="closed"):
        source.index.connection.execute("SELECT 1")


def test_limits_api_separates_xlsx_json_and_legacy_route(authenticated):
    normal = authenticated.get("/api/v1/acquisitions/limits", params={"format": "XLSX"})
    assert normal.status_code == 200, normal.text
    assert normal.json()["limits"]["data_rows"]["max"] == 1000000
    assert normal.json()["physical_sheet_rows"] == 1048576
    legacy = authenticated.get("/api/v1/acquisitions/limits", params={"format": "XLSX", "route": "LEGACY_UPLOAD"}).json()
    ordinary = authenticated.get("/api/v1/acquisitions/limits", params={"format": "JSON"}).json()
    assert legacy["limits"]["data_rows"]["max"] == ordinary["limits"]["data_rows"]["max"] == 100000
    assert legacy["limits"]["compressed_bytes"]["max"] == ordinary["limits"]["compressed_bytes"]["max"] == 10485760


def test_failed_row_limit_persists_same_safe_diagnostic_and_no_partial_version(authenticated, database, tmp_path, monkeypatch):
    fixture = fixtures.generate(tmp_path, 101, strings="shared")
    path = Path(fixture["path"])
    monkeypatch.setenv("TRACKVANCE_ACQUISITION_XLSX_MAX_ROWS", "100")
    monkeypatch.setenv("TRACKVANCE_ACQUISITION_BATCH_ROWS", "2")
    staged = authenticated.post("/api/v1/datasets/uploads/stage", params={"filename": path.name}, content=path.read_bytes())
    assert staged.status_code == 201, staged.text
    inspected = authenticated.get(f'/api/v1/datasets/uploads/{staged.json()["upload"]["id"]}/inspect',
                                  params={"reader_options": json.dumps({"sheet_name": "Datos"})})
    assert inspected.status_code == 200, inspected.text
    dataset = authenticated.post("/api/v1/datasets", json={"name": "XLSX with exact limit"}).json()
    registered = authenticated.post(f'/api/v1/datasets/{dataset["id"]}/acquisitions',
        json={"upload_id": staged.json()["upload"]["id"], "reader_options": {"sheet_name": "Datos"}})
    assert registered.status_code == 202, registered.text
    identity = registered.json()["id"]
    assert registered.json()["received_bytes"] == path.stat().st_size
    assert registered.json()["published_rows"] is None
    assert acquisition.process_once("xlsx-owner")
    detail = authenticated.get(f"/api/v1/acquisitions/{identity}").json()
    history = authenticated.get("/api/v1/acquisitions", params={"dataset_id": dataset["id"]}).json()["items"][0]
    assert detail["status"] == "FAILED" and detail["error"] == history["error"]
    assert detail["error"]["code"] == "ACQUISITION_ROW_LIMIT"
    assert detail["error"]["details"] == {"limit": "data_rows", "maximum": 100, "observed": 101}
    assert detail["error"]["reference"] == identity
    assert detail["materialized_rows"] > 0 and detail["published_rows"] is None
    with database() as db:
        run = db.get(AcquisitionRun, identity)
        assert run.output_version_id is None
        assert db.scalar(select(func.count()).select_from(DatasetVersion).where(DatasetVersion.dataset_id == dataset["id"])) == 0
        original = Path(db.get(AcquisitionUpload, run.upload_id).path)
        assert original.exists() and file_hash(original) == fixture["file_sha256"]
        event = db.scalar(select(OutboxEvent).where(OutboxEvent.aggregate_id == identity))
        assert event.payload["error_code"] == detail["error"]["code"]
        assert event.payload["error_message"] == detail["error"]["message"]
        assert event.payload["error_details"] == detail["error"]["details"]
        assert event.payload["error_reference"] == identity
        consume_notification(db, event)
        db.commit()
    inbox = authenticated.get("/api/v1/notifications/inbox").json()["items"]
    notification = next(item for item in inbox if item["resource_id"] == identity)
    assert notification["description"] == detail["error"]["message"]
    assert notification["error"] == detail["error"]


def test_success_reclaimed_xlsx_lease_publishes_full_profile_and_numbering(authenticated, database, tmp_path):
    fixture = fixtures.generate(tmp_path, 137, strings="shared")
    path = Path(fixture["path"])
    staged = authenticated.post("/api/v1/datasets/uploads/stage", params={"filename": path.name}, content=path.read_bytes()).json()
    dataset = authenticated.post("/api/v1/datasets", json={"name": "Recovered workbook"}).json()
    registered = authenticated.post(f'/api/v1/datasets/{dataset["id"]}/acquisitions',
        json={"upload_id": staged["upload"]["id"], "reader_options": {"sheet_name": "Datos"}})
    assert registered.status_code == 202, registered.text
    identity = registered.json()["id"]
    with database() as db:
        run = db.get(AcquisitionRun, identity)
        job = db.scalar(select(Job).where(Job.acquisition_id == identity))
        run.status, run.attempt_id, run.attempts = "RUNNING", "dead-attempt", 1
        job.status, job.attempts, job.lease_owner = "RUNNING", 1, "dead-worker"
        job.lease_until = utcnow() - timedelta(seconds=1)
        db.commit()
    assert acquisition.process_once("recovery-worker")
    with database() as db:
        run = db.get(AcquisitionRun, identity)
        assert run.status == "SUCCESS", (run.error_code, run.error_message)
        version = db.get(DatasetVersion, run.output_version_id)
        assert version.row_count == version.profile["row_count"] == 137
        assert version.ingestion_metadata["reader_version"] == 2
        rows = pl.concat(list(iter_version_batches(db, version, include_record_numbers=True)))
        assert rows["__tv_record_number"].to_list() == list(range(5, 142))
        digest = hashlib.sha256()
        for row in rows.drop("__tv_record_number").iter_rows(named=True):
            digest.update(fixtures.canonical(row))
        assert digest.hexdigest() == fixture["canonical_rows_sha256"]
        assert run.attempts == 2


def test_unrecognized_legacy_errors_do_not_leak_arbitrary_text():
    diagnostic = processing_diagnostic(ProcessingError("ACQUISITION_FAKE: /private/file credential=secret business-value"))
    assert diagnostic.code == "ACQUISITION_INVALID_DATA"
    assert "secret" not in diagnostic.message and "private" not in diagnostic.message
    row_limit = processing_diagnostic(ProcessingError("Este prototipo admite hasta 100,000 filas por archivo."))
    assert row_limit.code == "ACQUISITION_ROW_LIMIT" and row_limit.details["maximum"] == 100000


def test_sparse_physical_row_identity_preserves_header_offset_and_empty_formatted_gaps(tmp_path):
    path = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    sheet["A4"], sheet["A5"], sheet["A9"], sheet["A1000"] = "record_id", "first", "second", "third"
    sheet["A8"].number_format = "0.00"
    sheet["A1048576"].number_format = "0.00"
    workbook.save(path)
    changed = modify_package(path, tmp_path / "sparse-wrong-dimension.xlsx", lambda name, content:
        content.replace(b'<dimension ref="A4:A1048576"', b'<dimension ref="A1:A1"') if name.endswith("sheet1.xml") else content)
    reader = FileBatchReader(changed, changed.name, {"sheet_name": "Datos"}, replace(AcquisitionLimits(), batch_rows=1))
    batches = list(reader)
    assert [number for batch in batches for number in batch.record_numbers] == [5, 9, 1000]
    assert [batch.frame["record_id"][0] for batch in batches] == ["first", "second", "third"]
    assert reader.total_rows == 3
    assert all(batch.row_numbering == "PHYSICAL_SHEET_ROW" for batch in batches)
    inspection = inspect_file(changed, changed.name, {"sheet_name": "Datos"})
    assert inspection["row_numbering"] == "PHYSICAL_SHEET_ROW"
    assert inspection["effective_limits"]["header_row_number"] == 4


@pytest.mark.parametrize("lost_lease", [False, True])
def test_worker_cancel_or_lease_loss_during_shared_string_index_never_publishes(authenticated, database, tmp_path, monkeypatch, lost_lease):
    fixture = fixtures.generate(tmp_path, 1000, strings="shared")
    path = Path(fixture["path"])
    staged = authenticated.post("/api/v1/datasets/uploads/stage", params={"filename": path.name}, content=path.read_bytes()).json()
    dataset = authenticated.post("/api/v1/datasets", json={"name": "Cancel or fence indexing"}).json()
    registered = authenticated.post(f'/api/v1/datasets/{dataset["id"]}/acquisitions',
        json={"upload_id": staged["upload"]["id"], "reader_options": {"sheet_name": "Datos"}})
    assert registered.status_code == 202, registered.text
    identity, changed, put = registered.json()["id"], [], _Index.put
    def interfere(index, position, value):
        put(index, position, value)
        if position == 200 and not changed:
            with database() as db:
                if lost_lease:
                    job = db.scalar(select(Job).where(Job.acquisition_id == identity))
                    job.lease_until = utcnow() - timedelta(seconds=1)
                else:
                    db.get(AcquisitionRun, identity).cancel_requested = True
                db.commit()
            changed.append(True)
    monkeypatch.setattr(_Index, "put", interfere)
    assert acquisition.process_once("first-worker")
    assert changed == [True]
    with database() as db:
        run = db.get(AcquisitionRun, identity)
        assert run.output_version_id is None
        assert run.status == ("RUNNING" if lost_lease else "CANCELLED")
        assert run.processed_rows == 0
        assert db.scalar(select(func.count()).select_from(DatasetVersion).where(DatasetVersion.dataset_id == dataset["id"])) == 0
        assert file_hash(Path(db.get(AcquisitionUpload, run.upload_id).path)) == fixture["file_sha256"]
    if lost_lease:
        # The stale worker must not write a terminal state. A new claim gets a
        # distinct attempt and rereads the immutable original successfully.
        assert acquisition.process_once("reclaimed-worker")
        with database() as db:
            run = db.get(AcquisitionRun, identity)
            assert run.status == "SUCCESS", run.error_message
            assert run.attempts == 2 and run.processed_rows == 1000


def test_http_unknown_sheet_and_oversized_decimal_are_specific_and_sanitized(authenticated, tmp_path):
    fixture = fixtures.generate(tmp_path, 3)
    path = Path(fixture["path"])
    staged = authenticated.post("/api/v1/datasets/uploads/stage", params={"filename": path.name}, content=path.read_bytes()).json()
    response = authenticated.get(f'/api/v1/datasets/uploads/{staged["upload"]["id"]}/inspect',
        params={"reader_options": json.dumps({"sheet_name": "private-business-value"})})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "ACQUISITION_SHEET_NOT_FOUND"
    assert "private-business-value" not in response.text
    changed = modify_package(path, tmp_path / "expanded-decimal.xlsx", lambda name, content:
        content.replace(b"<v>2.00000001</v>", b"<v>1e200000</v>") if name.endswith("sheet2.xml") else content)
    with pytest.raises(AcquisitionReadError) as error:
        inspect_file(changed, changed.name, {"sheet_name": "Datos"})
    assert error.value.code == "ACQUISITION_CELL_LIMIT"
    assert error.value.details == {"limit": "cell_bytes", "maximum": 65536}


def test_profile_spill_limit_is_the_remaining_total_attempt_budget(authenticated, database, tmp_path, monkeypatch):
    fixture = fixtures.generate(tmp_path, 137, strings="shared")
    path = Path(fixture["path"])
    staged = authenticated.post("/api/v1/datasets/uploads/stage", params={"filename": path.name}, content=path.read_bytes()).json()
    dataset = authenticated.post("/api/v1/datasets", json={"name": "Total temporary budget"}).json()
    registered = authenticated.post(f'/api/v1/datasets/{dataset["id"]}/acquisitions',
        json={"upload_id": staged["upload"]["id"], "reader_options": {"sheet_name": "Datos"}})
    assert registered.status_code == 202, registered.text
    seen, real_profile = [], acquisition.profile_paths
    def capture(paths, **kwargs):
        cap = kwargs["temp_byte_limit"]
        physical = sum(item.stat().st_size for item in paths[0].parent.iterdir() if item.is_file())
        assert cap + physical == kwargs["limits"].xlsx_temp_bytes
        assert not (paths[0].parent / "xlsx-index.sqlite").exists()
        seen.append(cap)
        return real_profile(paths, **kwargs)
    monkeypatch.setattr(acquisition, "profile_paths", capture)
    assert acquisition.process_once("spill-budget-worker")
    assert len(seen) == 1
    with database() as db:
        assert db.get(AcquisitionRun, registered.json()["id"]).status == "SUCCESS"


def test_central_directory_budget_applies_before_zipfile_allocates_it(tmp_path, monkeypatch):
    fixture = fixtures.generate(tmp_path, 3)
    data = bytearray(Path(fixture["path"]).read_bytes())
    offset = data.rfind(b"PK\x05\x06")
    struct.pack_into("<L", data, offset + 12, 9 * 1024**2)
    changed = tmp_path / "oversized-central-directory.xlsx"
    changed.write_bytes(data)
    monkeypatch.setattr(zipfile, "ZipFile", lambda *_: pytest.fail("central directory budget must precede ZipFile construction"))
    reader = XlsxStreamingReader(changed, ReaderOptions(sheet_name="Datos"), AcquisitionLimits(), lambda: None)
    with pytest.raises(AcquisitionReadError) as error:
        list(reader)
    assert error.value.code == "ACQUISITION_XLSX_METADATA_LIMIT"
    assert error.value.details["maximum"] == 8 * 1024**2


def test_xlsx_macros_encrypted_members_duplicates_and_columns_are_rejected(tmp_path):
    fixture = fixtures.generate(tmp_path, 3)
    original = Path(fixture["path"])
    macro = tmp_path / "with-macro.xlsx"
    modify_package(original, macro, lambda _, content: content)
    with zipfile.ZipFile(macro, "a") as package:
        package.writestr("custom/vbaProject.bin", b"macro")
    with pytest.raises(AcquisitionReadError) as error:
        list(FileBatchReader(macro, macro.name, {"sheet_name": "Datos"}))
    assert error.value.code == "ACQUISITION_XLSX_UNSUPPORTED_CONTENT"
    duplicate = tmp_path / "duplicate-member.xlsx"
    modify_package(original, duplicate, lambda _, content: content)
    with pytest.warns(UserWarning, match="Duplicate name"), zipfile.ZipFile(duplicate, "a") as package:
        package.writestr("xl/workbook.xml", b"duplicate")
    with pytest.raises(AcquisitionReadError) as error:
        list(FileBatchReader(duplicate, duplicate.name, {"sheet_name": "Datos"}))
    assert error.value.code == "ACQUISITION_XLSX_INVALID_STRUCTURE"
    data = bytearray(original.read_bytes())
    local, central = data.find(b"PK\x03\x04"), data.find(b"PK\x01\x02")
    struct.pack_into("<H", data, local + 6, 1)
    struct.pack_into("<H", data, central + 8, 1)
    encrypted = tmp_path / "encrypted-member.xlsx"
    encrypted.write_bytes(data)
    with pytest.raises(AcquisitionReadError) as error:
        list(FileBatchReader(encrypted, encrypted.name, {"sheet_name": "Datos"}))
    assert error.value.code == "ACQUISITION_XLSX_INVALID_STRUCTURE"
    source = tmp_path / "many-columns.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append([f"column_{column}" for column in range(101)])
    worksheet.append(["value"] * 101)
    workbook.save(source)
    with pytest.raises(AcquisitionReadError) as error:
        list(FileBatchReader(source, source.name))
    assert error.value.code == "ACQUISITION_COLUMN_LIMIT"
    assert error.value.details == {"limit": "columns", "maximum": 100, "observed": 101}


@pytest.mark.parametrize("exception,code", [
    (MemoryError("secret payload /private/path"), "ACQUISITION_MEMORY_LIMIT"),
    (PermissionError("secret payload /private/path"), "ACQUISITION_STORAGE_PERMISSION"),
    (RuntimeError("secret payload /private/path"), "ACQUISITION_FAILED"),
])
def test_unknown_or_resource_worker_failure_preserves_reference_without_exception_text(authenticated, database, tmp_path, monkeypatch, caplog, exception, code):
    fixture = fixtures.generate(tmp_path, 3)
    path = Path(fixture["path"])
    staged = authenticated.post("/api/v1/datasets/uploads/stage", params={"filename": path.name}, content=path.read_bytes()).json()
    dataset = authenticated.post("/api/v1/datasets", json={"name": "Safe worker failure"}).json()
    registered = authenticated.post(f'/api/v1/datasets/{dataset["id"]}/acquisitions',
        json={"upload_id": staged["upload"]["id"], "reader_options": {"sheet_name": "Datos"}})
    assert registered.status_code == 202, registered.text
    def fail(*_):
        raise exception
    monkeypatch.setattr(acquisition, "execute_acquisition", fail)
    assert acquisition.process_once("safe-error-worker")
    result = authenticated.get(f'/api/v1/acquisitions/{registered.json()["id"]}')
    assert result.json()["error"]["code"] == code
    assert result.json()["error"]["reference"] == registered.json()["id"]
    assert "secret" not in result.text and "private" not in result.text
    assert "secret" not in caplog.text and "private" not in caplog.text
    assert result.json()["received_bytes"] > 0
    assert result.json()["materialized_rows"] == 0 and result.json()["published_rows"] is None


def test_optional_cell_coordinates_keep_shared_formula_origin_and_physical_rows(tmp_path):
    source = tmp_path / "coordinates.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    for column, value in enumerate(("record_id", "amount", "formula"), 1):
        sheet.cell(4, column, value)
    for physical, identifier, amount in ((5, "first", 2), (9, "second", 3)):
        sheet.cell(physical, 1, identifier)
        sheet.cell(physical, 2, amount)
        sheet.cell(physical, 3, f"=A{physical}+B{physical}")
    workbook.save(source)
    def transform(name, content):
        if not name.endswith("sheet1.xml"):
            return content
        content = content.replace(b"<f>A5+B5</f>", b'<f t="shared" si="0">A5+B5</f>')
        content = content.replace(b"<f>A9+B9</f>", b'<f t="shared" si="0"/>')
        return re.sub(br'(<c) r="[^"]+"', br"\1", content)
    changed = modify_package(source, tmp_path / "implicit-cell-coordinates.xlsx", transform)
    batches = list(FileBatchReader(changed, changed.name, {"sheet_name": "Datos"}))
    assert [number for batch in batches for number in batch.record_numbers] == [5, 9]
    assert pl.concat([batch.frame for batch in batches]).to_dicts() == [
        {"record_id": "first", "amount": "2", "formula": "=A5+B5"},
        {"record_id": "second", "amount": "3", "formula": "=A9+B9"},
    ]
    no_row_coordinate = modify_package(source, tmp_path / "implicit-row-coordinate.xlsx", lambda name, content:
        content.replace(b'<row r="9">', b"<row>") if name.endswith("sheet1.xml") else content)
    assert [number for batch in FileBatchReader(no_row_coordinate, no_row_coordinate.name) for number in batch.record_numbers] == [5, 9]


def test_historical_xlsx_does_not_claim_new_limits_or_invent_error_details():
    frozen = {name: value for name, value in AcquisitionLimits().as_dict().items() if not name.startswith("xlsx_")}
    run = AcquisitionRun(id="historical-acquisition", dataset_id="historical-dataset", source_type="UPLOAD",
        filename="historical.xlsx", source_snapshot={}, reader_options={}, column_overrides={}, effective_limits=frozen,
        status="FAILED", stage="FAILED", attempts=1, attempt_id="old-attempt", initiated_by_id="old-user",
        initiated_by_name="Old user", processed_rows=0, processed_bytes=0, total_bytes=12345, total_rows=None,
        cancel_requested=False, output_version_id=None, error_code="ACQUISITION_INVALID_DATA",
        error_message="Mensaje histórico sin causa clasificada.", error_details=None, error_reference=None,
        created_at=utcnow(), started_at=None, finished_at=utcnow())
    before = json.dumps(frozen, sort_keys=True)
    result = acquisition.acquisition_dto(run)
    assert result["route_limits"] is None
    assert result["error"] == {"code": "ACQUISITION_INVALID_DATA", "message": run.error_message,
                               "details": None, "reference": None}
    assert json.dumps(run.effective_limits, sort_keys=True) == before


def test_iso_xlsx_dates_preserve_timezone_fraction_and_naive_serial_is_not_invented_timestamp(tmp_path):
    source = tmp_path / "observed-temporals.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    columns = ["aware_z", "aware_offset", "naive_iso", "nanoseconds_iso", "date_iso", "serial_datetime"]
    sheet.append(columns)
    sheet.append(["placeholder"] * 5 + [datetime.fromisoformat("2026-01-01T12:34:56.123000")])
    workbook.save(source)
    exact = ["2026-01-01T12:34:56.123456Z", "2026-01-01T12:34:56.123456-05:00", "2026-01-01T12:34:56.123456",
             "2026-01-01T12:34:56.123456789Z", date(2026, 1, 1).isoformat()]
    def transform(name, content):
        if not name.endswith("sheet1.xml"):
            return content
        for column, value in zip("ABCDE", exact, strict=True):
            content = re.sub(fr'<c r="{column}2".*?</c>'.encode(), f'<c r="{column}2" t="d"><v>{value}</v></c>'.encode(), content)
        return content
    changed = modify_package(source, tmp_path / "temporal-iso-values.xlsx", transform)
    reader = FileBatchReader(changed, changed.name)
    batches = list(reader)
    observed = pl.concat([batch.frame for batch in batches]).row(0)
    assert observed == (*exact, "2026-01-01T12:34:56.123000")
    path = tmp_path / "temporal-population.parquet"
    batches[0].physical_frame().write_parquet(path)
    schema, profile, _ = profile_paths([path], native_types=reader.native_schema)
    assert [column["logical_type"] for column in schema] == ["TIMESTAMP", "TIMESTAMP", "STRING", "STRING", "DATE", "STRING"]
    assert all(column.get("parse_error_count", 0) == 0 for column in profile["columns"])


@pytest.mark.parametrize("strings", ["inline", "shared"])
def test_empty_preheader_and_trailing_header_cells_preserve_data_and_physical_rows(tmp_path, strings):
    source = tmp_path / "header-empty.xlsx"
    workbook = Workbook()
    workbook.active.append(["id", "amount"])
    workbook.save(source)
    texts = ["", "id", "amount", "   ", "007"]
    def cell(column, physical, value):
        if strings == "shared":
            return f'<c r="{column}{physical}" t="s"><v>{texts.index(value)}</v></c>'
        return f'<c r="{column}{physical}" t="inlineStr"><is><t xml:space="preserve">{value}</t></is></c>'
    xml = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
           f'<row r="1">{cell("A", 1, "")}</row><row r="2"><c r="A2" s="0"/></row>'
           f'<row r="4">{cell("A", 4, "id")}{cell("B", 4, "amount")}{cell("C", 4, "")}{cell("D", 4, "   ")}</row>'
           f'<row r="5">{cell("A", 5, "")}{cell("B", 5, "   ")}</row>'
           f'<row r="7">{cell("A", 7, "007")}<c r="B7"/></row>'
           f'<row r="9"><c r="A9"/>{cell("B", 9, "")}</row>'
           '</sheetData></worksheet>').encode()
    changed = tmp_path / f"empty-{strings}.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(changed, "x", compression=zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            target.writestr(item.filename, xml if item.filename.endswith("sheet1.xml") else original.read(item.filename))
        if strings == "shared":
            sst = '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">' + ''.join(
                f'<si><t xml:space="preserve">{value}</t></si>' for value in texts) + '</sst>'
            target.writestr("xl/sharedStrings.xml", sst)
    expected = [{"id": "", "amount": "   "}, {"id": "007", "amount": None}, {"id": None, "amount": ""}]
    inspection = inspect_file(changed, changed.name)
    assert [column["name"] for column in inspection["columns"]] == ["id", "amount"]
    assert inspection["effective_limits"]["header_row_number"] == 4
    assert inspection["sample"] == expected
    reader = FileBatchReader(changed, changed.name)
    batches = list(reader)
    assert pl.concat([batch.frame for batch in batches]).to_dicts() == expected
    assert [number for batch in batches for number in batch.record_numbers] == [5, 7, 9]
    assert reader.metadata["physical_header_row_number"] == 4


@pytest.mark.parametrize("strings", ["inline", "shared"])
def test_auto_inspection_skips_definitively_empty_sheet_and_matches_worker(tmp_path, strings):
    original = tmp_path / "empty-first.xlsx"
    workbook = Workbook()
    workbook.active.title = "Vacía"
    workbook.create_sheet("Datos")
    workbook.save(original)
    texts = ["", "id", "value", "007", "   "]
    def cell(column, physical, value):
        body = f'<v>{texts.index(value)}</v>' if strings == "shared" else f'<is><t xml:space="preserve">{value}</t></is>'
        kind = "s" if strings == "shared" else "inlineStr"
        return f'<c r="{column}{physical}" t="{kind}">{body}</c>'
    prefix = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
    sheets = {"xl/worksheets/sheet1.xml": (prefix + f'<row r="1">{cell("A", 1, "")}</row></sheetData></worksheet>').encode(),
              "xl/worksheets/sheet2.xml": (prefix + f'<row r="4">{cell("A", 4, "id")}{cell("B", 4, "value")}</row>'
                                          f'<row r="5">{cell("A", 5, "007")}{cell("B", 5, "   ")}</row></sheetData></worksheet>').encode()}
    changed = tmp_path / f"auto-{strings}.xlsx"
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(changed, "x", compression=zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, sheets.get(item.filename, source.read(item.filename)))
        if strings == "shared":
            target.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">' + ''.join(
                f'<si><t xml:space="preserve">{value}</t></si>' for value in texts) + '</sst>')
    expected = [{"id": "007", "value": "   "}]
    for options in [None, {"sheet_name": "Datos"}]:
        inspection = inspect_file(changed, changed.name, options)
        assert inspection["selected_sheet"] == "Datos" and inspection["sample"] == expected
        assert inspection["effective_limits"]["header_row_number"] == 4
        reader = FileBatchReader(changed, changed.name, options)
        batch = next(iter(reader))
        assert batch.frame.to_dicts() == expected and batch.record_numbers == [5]
