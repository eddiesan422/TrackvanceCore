import copy
import io
import zipfile
from datetime import datetime

import pytest
from openpyxl import load_workbook

from trackvance.exports import build_run_workbook, export_filename


def example(module="intake", rows=None):
    config = {"id": "config-1", "version": 2, "name": "Control Colombia", "config": {"max_error_rate": .05, "key_columns": ["document_id"], "amount_column": "amount", "tolerance": "0.01"}}
    run = {"id": "run-1", "module": module, "status": "SUCCESS", "decision": "REJECTED" if module == "intake" else "WITH_FINDINGS" if module == "recon" else "ALERT", "started_at": "2026-09-01T10:00:00Z", "finished_at": "2026-09-01T10:01:00Z", "metrics": {"total_rows": 120, "valid_rows": 110, "error_rows": 10, "warning_rows": 0, "acceptance_rate": 91.67, "source_rows": 120, "target_rows": 110, "counts": {"MATCH": 110, "SOURCE_ONLY": 6, "DUPLICATE_SOURCE": 4, "VALUE_MISMATCH": 0, "TARGET_ONLY": 0, "DUPLICATE_TARGET": 0, "INVALID": 0}, "match_rate": 91.67, "row_count": 90, "total_checks": 9, "failed_checks": 1, "health_score": 88.89, "rules": [{"code": "REQUIRED", "column": "email", "evaluated_count": 120, "failed_count": 4, "status": "FAIL"}]}}
    inputs = [{"id": "dv-1", "dataset_id": "ds-1", "dataset_name": "Ventas", "version": 1, "sha256": "a" * 64, "schema_hash": "b" * 64, "source_type": "UPLOAD", "original_artifact_id": "original-1", "canonical_artifact_id": "canonical-1"}]
    if module == "recon":
        inputs.append({**inputs[0], "id": "dv-2", "dataset_name": "Pagos"})
    manifest = {"schema_version": 2, "run_id": "run-1", "initiated_by": {"type": "USER", "id": "user-1", "display_name": "Equipo Trackvance"}, "configuration": {**config, "config_hash": "c" * 64}, "processing": {"engine": "POLARS", "engine_version": "test"}, "result_artifacts": [{"artifact_id": "results-1", "kind": "RECON_RESULTS", "name": "results.parquet", "sha256": "d" * 64, "size_bytes": 1024}]}
    return run, config, inputs, manifest, rows or []


@pytest.mark.parametrize("value", ["=1+1", "+CMD|'/C calc'!A0", "-1+2", "@SUM(A1:A2)", "\t=HYPERLINK(\"https://example.test\")", "  =1+1", "001234567", "é", ""])
def test_received_values_are_literal_text(value):
    row = {"original_row_number": 2, "rule_code": "ALLOWED_VALUES", "column": "document_id", "received_value": value, "severity": "ERROR", "message": "Valor no permitido", "classification": "ERROR"}
    content = build_run_workbook(*example(rows=[row]))
    book = load_workbook(io.BytesIO(content))
    cell = book["Errores"]["D6"]
    assert cell.data_type != "f" and (cell.value == value or (value == "" and cell.value is None))
    assert cell.number_format == "@"
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        assert b"<f>" not in z.read("xl/worksheets/sheet2.xml")


@pytest.mark.parametrize("module,sheets", [("intake", ["Resumen", "Errores", "Reglas", "Trazabilidad"]), ("recon", ["Resumen", "Resultados", "Hallazgos", "Trazabilidad"]), ("sentinel", ["Resumen", "Controles", "Hallazgos", "Trazabilidad"])])
def test_workbook_structure_identity_and_typed_metrics(module, sheets):
    arguments = example(module)
    before = copy.deepcopy(arguments)
    book = load_workbook(io.BytesIO(build_run_workbook(*arguments)))
    assert arguments == before
    assert book.sheetnames == sheets
    for sheet in book:
        assert sheet.freeze_panes and sheet.tables and sheet.auto_filter.ref
        assert sheet.sheet_view.showGridLines is False
        assert sheet["A5"].fill.fgColor.rgb.endswith("172D4D")
    assert isinstance(book["Resumen"]["B12"].value, datetime)
    assert any(c.number_format == "0.00%" and isinstance(c.value, float) for row in book["Resumen"] for c in row)
    trace = str(list(book["Trazabilidad"].values))
    assert "results-1" in trace and "d" * 64 in trace and "user-1" in trace


def test_recon_findings_exclude_matches_and_keep_decimal_values():
    rows = [{"key": "001234567", "classification": kind, "source_value": "100.25", "target_value": "101.50", "difference": "-1.25", "tolerance": "0.01", "source_row": 2, "target_row": 3, "message": "Diferencia", "comparisons": [{"code": "NUMERIC_TOLERANCE", "status": "FAIL"}]} for kind in ["MATCH", "VALUE_MISMATCH"]]
    book = load_workbook(io.BytesIO(build_run_workbook(*example("recon", rows))))
    assert book["Resultados"]["A6"].value == "001234567"
    assert book["Resultados"]["C6"].value == 100.25
    assert book["Resultados"]["E6"].value == -1.25
    assert "MATCH)" not in str(list(book["Hallazgos"].values)).replace("VALUE_MISMATCH)", "")
    assert len(book["Resultados"].conditional_formatting) > 0


def test_exports_exclude_secret_metadata_and_sanitize_filename():
    run, config, inputs, manifest, rows = example()
    manifest["configuration"]["config"].update(password="DONOTEXPORT", nested={"api_token": "DONOTEXPORT"})
    book = load_workbook(io.BytesIO(build_run_workbook(run, config, inputs, manifest, rows)))
    assert "DONOTEXPORT" not in str([list(sheet.values) for sheet in book])
    name = export_filename('../intake\r\n', '../../run"\r\nX-Evil:1')
    assert '/' not in name and '\\' not in name and '\r' not in name and '"' not in name and name.endswith('.xlsx')


def test_recon_export_uses_actual_boolean_check_and_readable_configuration():
    arguments = example("recon", [{"key": "A", "classification": "MATCH", "source_value": "10", "target_value": "10", "comparisons": [{"passed": True, "code": "NUMERIC_TOLERANCE", "source_column": "amount", "target_column": "amount", "source_value": "10", "target_value": "10"}]}])
    book = load_workbook(io.BytesIO(build_run_workbook(*arguments)))
    assert "Cumple (PASS)" in book["Resultados"]["J6"].value
    summary = str(list(book["Resumen"].values))
    assert "Absoluta: 0.01" in summary and "Normalización de claves" in summary
    assert "schema_version" not in summary


def test_openapi_exposes_xlsx_cookie_and_csrf_contract():
    from trackvance.api import app

    schema = app.openapi()
    assert schema["info"]["version"] == "0.3.0"
    export = schema["paths"]["/api/v1/runs/{run_id}/export.xlsx"]["get"]
    assert set(export["responses"]["200"]["content"]) == {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    assert export["security"] == [{"LocalSession": []}]
    assert schema["components"]["securitySchemes"]["LocalSession"]["in"] == "cookie"
    mutation = schema["paths"]["/api/v1/intake/contracts"]["post"]
    assert any(p["name"] == "X-CSRF-Token" and p["required"] for p in mutation["parameters"])
