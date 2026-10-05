"""Independent semantic oracles for AST authorization and no-spool exports."""
import csv
import io
import zipfile
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from trackvance.operations_common import OperationError
from trackvance.report_config import ReportLimits
from trackvance.report_exports import csv_stream, preview_row, xlsx_stream
from trackvance.report_query import compile_draft, typed_parameters, validate_sql

SCHEMA = {"a": [{"name": "id", "logical_type": "STRING"}, {"name": "amount", "logical_type": "DECIMAL"}],
          "b": [{"name": "key", "logical_type": "STRING"}, {"name": "value", "logical_type": "INT64"}]}
POLICY = [{"expected_cardinality": "1:N", "allow_many_to_many": False}]


@pytest.mark.parametrize("sql", [
    "SELECT a.id FROM a; SELECT b.key FROM b",
    "DELETE FROM a", "COPY a TO '/tmp/escape.csv'",
    "SELECT a.id FROM a UNION ALL SELECT b.key FROM b",
    "SELECT a.id FROM a CROSS JOIN b", "SELECT a.id FROM a NATURAL JOIN b",
    "SELECT a.id FROM a JOIN b ON 1=1", "SELECT a.id FROM a JOIN b ON a.id=a.id",
    "SELECT a.id FROM a JOIN b ON a.id=b.key OR a.id=b.key",
    "SELECT a.id FROM a JOIN b ON a.id>b.key",
    "SELECT a.id FROM a, b", "SELECT * FROM a JOIN b ON a.id=b.key",
    "SELECT a.id FROM a JOIN read_parquet('/tmp/secret') b ON a.id=b.key",
    "WITH hidden AS (SELECT * FROM read_csv('https://example.test')) SELECT a.id FROM a JOIN b ON a.id=b.key",
    "SELECT a.id,(SELECT count(*) FROM duckdb_settings()) AS hidden FROM a JOIN b ON a.id=b.key",
    "SELECT random() AS unsafe FROM a JOIN b ON a.id=b.key",
    "SELECT current_timestamp AS unsafe FROM a JOIN b ON a.id=b.key",
    "SELECT getenv('DATABASE_URL') AS unsafe FROM a JOIN b ON a.id=b.key",
    "SELECT a.missing FROM a JOIN b ON a.id=b.key",
    "SELECT a.id FROM main.a JOIN b ON a.id=b.key",
    "SELECT a.id AS same,b.key AS same FROM a JOIN b ON a.id=b.key",
])
def test_ast_rejects_entire_unauthorized_tree(sql):
    with pytest.raises(OperationError):
        validate_sql(sql, SCHEMA, POLICY, {})


def test_guided_values_are_bound_and_source_filters_stay_separate():
    malicious = "' OR 1=1 --"
    draft = {"sources": [{"alias": "a"}, {"alias": "b"}], "mode": "GUIDED",
             "joins": [{"left_alias": "a", "right_alias": "b", "type": "LEFT",
                        "keys": [{"left_column": "id", "right_column": "key"}], **POLICY[0]}],
             "columns": [{"source_alias": "a", "column": "id", "alias": "identifier"}],
             "source_filters": {"a": {"column": "id", "operator": "EQ", "value": malicious}},
             "post_filter": {"source_alias": "b", "column": "value", "operator": "IS_NULL"}}
    result = compile_draft(draft, SCHEMA)
    assert malicious not in result["sql"]
    assert list(result["parameters"].values()) == [malicious]
    assert '"b"."value" IS NULL' in result["sql"]
    assert '"a"."id"' in result["source_filters"]["a"]
    assert "__tv_source_pos" in result["sql"]


def test_sql_aggregate_and_named_parameters_preserve_contract():
    plan = validate_sql("SELECT a.id AS identifier,SUM(a.amount) AS exact_amount,COUNT(*) AS records FROM a LEFT JOIN b ON a.id=b.key WHERE b.value>$threshold GROUP BY a.id HAVING COUNT(*)>0 ORDER BY identifier",
                        SCHEMA, POLICY, {"threshold": 12})
    assert plan["columns"] == ["identifier", "exact_amount", "records"]
    assert "$threshold" in plan["sql"]
    assert plan["used_columns"]["b"] == ["key", "value"]
    with pytest.raises(OperationError, match="AVG"):
        validate_sql("SELECT AVG(a.amount) AS mean FROM a JOIN b ON a.id=b.key", SCHEMA, POLICY, {})


def test_int64_parameter_does_not_pass_through_js_float():
    assert typed_parameters([{"name": "exact", "type": "INTEGER", "value": "9223372036854775807"}]) == {"exact": 9223372036854775807}
    with pytest.raises(OperationError):
        typed_parameters([{"name": "exact", "type": "INTEGER", "value": "9223372036854775808"}])
    assert preview_row(["exact", "ordinary", "boolean"], [9223372036854775807, 12, True]) == {
        "exact": "9223372036854775807", "ordinary": 12, "boolean": True}


@pytest.mark.parametrize("value", ["1e38", "1e-39", "NaN", "Infinity"])
def test_decimal_parameters_cannot_silently_become_float_or_lose_scale(value):
    with pytest.raises(OperationError):
        typed_parameters([{"name": "exact", "type": "DECIMAL", "value": value}])


def test_preview_cannot_override_ten_row_cap(monkeypatch):
    monkeypatch.setenv("REPORT_PREVIEW_MAX_ROWS", "1000")
    assert ReportLimits.configured().max_rows == 10


def test_csv_null_empty_identifiers_unicode_and_formula_contract():
    columns = [{"name": n} for n in ["id", "amount", "nothing", "empty", "formula", "unicode"]]
    batch = [["0001", {"type": "DECIMAL", "value": "12345678901234567890.001"}, None, "", "=1+1", "á東京"]]
    content = b"".join(csv_stream(columns, iter([batch]))).decode()
    values = list(csv.reader(io.StringIO(content)))[1]
    assert values == ["0001", "12345678901234567890.001", r"\N", "", "'=1+1", "á東京"]


def test_csv_formula_header_and_escaped_prefixes_are_reversible():
    values = [None, "", r"\N", r"\\literal", "'literal", " \t=1+1", "-001", "@formula", "東京"]
    content = b"".join(csv_stream([{"name": "=unsafe_header"}], iter([[[value] for value in values]]))).decode()
    rows = list(csv.reader(io.StringIO(content)))
    assert rows[0] == ["'=unsafe_header"]

    def decode(value):
        if value == r"\N":
            return None
        if value.startswith(("\\\\", "'")):
            return value[1:]
        return value

    assert [decode(row[0]) for row in rows[1:]] == values


def test_xlsx_never_calls_temporary_file_and_has_exact_text_cells():
    columns = [{"name": n} for n in ["id", "amount", "nothing", "empty", "formula", "bigint"]]
    values = [["0001", {"type": "DECIMAL", "value": "12345678901234567890.001"}, None, "", "=1+1", 9223372036854775807]]
    with patch("tempfile.TemporaryFile", side_effect=AssertionError("spool")), patch("tempfile.NamedTemporaryFile", side_effect=AssertionError("spool")):
        content = b"".join(xlsx_stream(columns, iter([values])))
    sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=False).active
    cells = list(sheet.iter_rows())[1]
    assert [cell.value for cell in cells] == ["0001", "12345678901234567890.001", None, "", "=1+1", "9223372036854775807"]
    assert cells[4].data_type == "s"
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        xml = package.read("xl/worksheets/sheet1.xml").decode()
    assert '<c r="C2"' not in xml
    assert '<c r="D2" t="inlineStr"><is><t xml:space="preserve"></t>' in xml


def test_xlsx_rejects_unrepresentable_xml_and_cell_length():
    for value in ["bad\x00", "x" * 32768]:
        with pytest.raises(OperationError):
            b"".join(xlsx_stream([{"name": "text"}], iter([[[value]]])))
