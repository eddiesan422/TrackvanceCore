from datetime import datetime
from io import BytesIO

import polars as pl
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from trackvance.artifactstore import (
    ArtifactIntegrityError,
    FileArtifactStore,
    StorageProvider,
)
from trackvance.dataset_readers import (
    DatasetReadResult,
    DatasetSource,
    LocalFileDatasetSource,
    inspect_dataset,
    read_dataset,
)
from trackvance.execution import ExecutionEngine
from trackvance.jobqueue import JobQueue
from trackvance.models import Artifact, Configuration, Dataset, DatasetVersion, Job, Run
from trackvance.services import create_version, enqueue
from trackvance.worker import process_once


def test_local_storage_adapter_implements_port_and_owns_paths(tmp_path):
    root = tmp_path / "provider"
    root.mkdir()
    provider = FileArtifactStore(root)

    assert isinstance(provider, StorageProvider)
    directory = provider.create_directory("runs", "run-1")
    assert directory == (root / "runs" / "run-1").resolve()
    assert directory.is_dir()
    staged = provider.temporary_path(".parquet")
    staged.write_bytes(b"parquet-test")
    assert staged.parent == (root / "tmp").resolve()
    assert provider.materialize_reference(str(staged)) == staged
    with pytest.raises(ArtifactIntegrityError, match="fuera"):
        provider.location("..", "outside")


def test_dataset_source_port_normalizes_non_file_connectors():
    class MemoryDatasetSource:
        source_kind = "MEMORY_TEST"

        def __init__(self):
            self.inspections: list[bool] = []

        def read(self, options=None, *, inspect=False):
            self.inspections.append(inspect)
            return DatasetReadResult(
                frame=pl.DataFrame({"id": ["001"], "amount": ["10.00"]}),
                source_format="MEMORY",
                format_label="Memory test",
                media_type="application/x-test",
                row_numbering="RECORD_NUMBER",
            )

    source = MemoryDatasetSource()
    assert isinstance(source, DatasetSource)
    assert read_dataset(source).frame.to_dicts() == [{"id": "001", "amount": "10.00"}]
    assert inspect_dataset(source).source_format == "MEMORY"
    assert source.inspections == [False, True]


def test_local_file_source_is_the_compatibility_adapter(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text("id,value\n001,10\n", encoding="utf-8")
    source = LocalFileDatasetSource(path, path.name)

    assert isinstance(source, DatasetSource)
    assert read_dataset(source).frame.to_dicts() == [{"id": "001", "value": "10"}]


def test_enqueue_accepts_a_queue_port_without_changing_run_semantics(database, tmp_path):
    class RecordingQueue:
        key = "TEST_QUEUE"

        def __init__(self):
            self.run_ids: list[str] = []

        def submit(self, db, run, *, executable=True):
            self.run_ids.append(run.id)
            job = Job(
                organization_id=run.organization_id,
                run_id=run.id,
                status="QUEUED" if executable else "FAILED",
            )
            db.add(job)
            return job

    source_path = tmp_path / "queue.csv"
    source_path.write_text("id\n001\n", encoding="utf-8")
    queue = RecordingQueue()
    assert isinstance(queue, JobQueue)

    with database() as db:
        dataset = Dataset(name="Queue port")
        db.add(dataset)
        db.flush()
        version = create_version(db, dataset, source_path, source_path.name)
        configuration = Configuration(
            name="Queue contract",
            module="intake",
            dataset_id=dataset.id,
            config={"required_columns": ["id"]},
        )
        db.add(configuration)
        db.flush()

        run = enqueue(db, configuration, version, None, "Test User", queue=queue)
        db.flush()

        assert queue.run_ids == [run.id]
        assert db.query(Job).filter_by(run_id=run.id, status="QUEUED").one()


def test_worker_delegates_complete_run_to_execution_engine_port(database, queued_intake):
    class RecordingEngine:
        key = "TEST_ENGINE"
        supported_processing_engines = frozenset({"POLARS"})

        def __init__(self):
            self.calls: list[tuple[str, str | None]] = []

        def execute(
            self,
            db,
            run,
            *,
            lease_owner=None,
            observed_at: datetime | None = None,
        ):
            self.calls.append((run.id, lease_owner))
            run.status = "SUCCESS"
            run.decision = "APPROVED"
            run.progress_percent = 100
            run.progress_stage = "Completado"

    engine = RecordingEngine()
    assert isinstance(engine, ExecutionEngine)
    assert process_once("port-worker", engine=engine)
    assert engine.calls == [(queued_intake["run_id"], "port-worker")]

    with database() as db:
        run = db.get(Run, queued_intake["run_id"])
        job = db.query(Job).filter_by(run_id=run.id).one()
        assert run.status == "SUCCESS"
        assert job.status == "SUCCESS"


def test_application_flows_materialize_opaque_storage_references(
    authenticated, database, monkeypatch, tmp_path
):
    """A provider locator is not a filesystem path, even with local compute."""
    from trackvance import api, services

    class OpaqueStorageProvider:
        def __init__(self):
            self.local = FileArtifactStore(tmp_path / "objects")
            self.references = {}

        def temporary_path(self, suffix=".tmp"):
            return self.local.temporary_path(suffix)

        def put_file(self, db, source, kind, organization_id, name=None,
                     artifact_id=None, media_type=None):
            artifact = self.local.put_file(
                db, source, kind, organization_id, name, artifact_id, media_type
            )
            reference = f"test-object://bucket/{artifact.id}"
            self.references[reference] = artifact.path
            artifact.path = reference
            db.flush()
            return artifact

        def materialize_reference(self, reference, *, expected_sha256=None, expected_size=None):
            return self.local.materialize_reference(
                self.references[reference],
                expected_sha256=expected_sha256,
                expected_size=expected_size,
            )

        def materialize(self, artifact):
            return self.materialize_reference(
                artifact.path,
                expected_sha256=artifact.sha256,
                expected_size=artifact.size_bytes,
            )

        def exists(self, artifact):
            return artifact.path in self.references

        def open_read(self, artifact):
            return self.materialize(artifact).open("rb")

    provider = OpaqueStorageProvider()
    assert isinstance(provider, StorageProvider)
    monkeypatch.setattr(api, "storage_provider", provider)
    monkeypatch.setattr(services, "storage_provider", provider)
    content = b"order_id,amount\n001,10\n002,-1\n"
    dataset = authenticated.post("/api/v1/datasets", json={"name": "Opaque storage"}).json()
    upload = authenticated.post(
        f"/api/v1/datasets/{dataset['id']}/versions/upload",
        files={"file": ("orders.csv", content, "text/csv")},
    )
    assert upload.status_code == 201
    version = upload.json()
    profile = authenticated.get(f"/api/v1/dataset-versions/{version['id']}/profile")
    assert profile.status_code == 200
    assert profile.json()["sample"] == [
        {"order_id": "001", "amount": "10"}, {"order_id": "002", "amount": "-1"}
    ]
    schema = authenticated.get(f"/api/v1/datasets/{dataset['id']}/schema?refresh=true")
    assert schema.status_code == 200
    assert [column["name"] for column in schema.json()["columns"]] == ["order_id", "amount"]
    original = authenticated.get(
        f"/api/v1/artifacts/{version['original_artifact_id']}/download"
    )
    assert original.status_code == 200 and original.content == content

    with database() as db:
        configuration = Configuration(
            name="Opaque storage contract", module="intake", dataset_id=dataset["id"],
            config={"positive_columns": ["amount"], "max_error_rate": 0},
        )
        db.add(configuration)
        db.flush()
        run = enqueue(
            db, configuration, db.get(DatasetVersion, version["id"]), None, "Test User"
        )
        run_id = run.id
        db.commit()

    assert process_once("opaque-storage-worker")
    detail = authenticated.get(f"/api/v1/runs/{run_id}").json()
    assert detail["status"] == "SUCCESS" and detail["decision"] == "REJECTED"
    assert authenticated.get(f"/api/v1/runs/{run_id}/results").json()["total"] == 1
    manifest = authenticated.get(f"/api/v1/runs/{run_id}/evidence")
    assert manifest.status_code == 200 and manifest.json()["run_id"] == run_id
    output = authenticated.get(
        f"/api/v1/dataset-versions/{detail['output_version_id']}/profile"
    )
    assert output.status_code == 200
    assert output.json()["sample"] == [{"order_id": "001", "amount": "10"}]

    exported = authenticated.get(f"/api/v1/runs/{run_id}/export.xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == ["Resumen", "Errores", "Reglas", "Trazabilidad"]
    workbook.close()
    with database() as db:
        artifacts = db.scalars(select(Artifact)).all()
        assert len(artifacts) >= 6
        assert all(artifact.path.startswith("test-object://") for artifact in artifacts)
