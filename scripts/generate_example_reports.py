"""Generate evidence using synthetic certified inputs through the running local API."""

import argparse
import importlib.util
import io
import json
import time
from datetime import UTC, datetime
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("certified_data", ROOT / "backend/tests/certified_data.py")
assert spec and spec.loader
fixtures = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixtures)


def generate(base_url: str, output: Path) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    with httpx.Client(base_url=base_url.rstrip('/') + '/api/v1', timeout=45) as client:
        login = client.post('/auth/demo', json={})
        login.raise_for_status()
        client.headers['X-CSRF-Token'] = login.json()['csrf_token']

        def post(path, **kwargs):
            response = client.post(path, **kwargs)
            response.raise_for_status()
            return response.json()

        def execute(path, data):
            queued = post(path, json=data)
            deadline = time.monotonic() + 90
            while time.monotonic() < deadline:
                response = client.get(f"/runs/{queued['id']}")
                response.raise_for_status()
                run = response.json()
                if run['status'] == 'SUCCESS':
                    return run
                if run['status'] in {'FAILED', 'FAILED_PRECONDITION', 'CANCELLED'}:
                    raise RuntimeError(f"El ejemplo no terminó correctamente: {run['status']} {run.get('error')}")
                time.sleep(.5)
            raise TimeoutError('La ejecución del ejemplo superó 90 segundos.')

        def upload(dataset, rows, columns=None):
            return post(f'/datasets/{dataset}/versions/upload', files={'file': ('certificacion.csv', fixtures.csv_bytes(rows, columns), 'text/csv')})

        stamp = datetime.now(UTC).strftime('%Y%m%dT%H%M%SZ')
        dataset = post('/datasets', json={'name': f'Informe de ejemplo {stamp}', 'description': 'Datos sintéticos: ciclo certificado 120/110/90 filas.'})
        original = fixtures.original_rows()
        version = upload(dataset['id'], original)
        contract = post('/intake/contracts', json={'name': f'Contrato certificado {stamp}', 'dataset_id': dataset['id'], 'config': fixtures.INTAKE_CONFIG})
        intake = execute('/intake/runs', {'contract_id': contract['id'], 'dataset_version_id': version['id']})
        assert intake['metrics']['valid_rows'] == 110 and intake['metrics']['error_rows'] == 10
        accepted_response = client.get(f"/dataset-versions/{intake['output_version_id']}/profile")
        accepted_response.raise_for_status()
        accepted = accepted_response.json()
        control = post('/recon/controls', json={'name': f'Conciliación certificada {stamp}', 'dataset_id': dataset['id'], 'target_dataset_id': accepted['dataset_id'], 'config': {'key_columns': ['transaction_id'], 'amount_column': 'total_amount', 'tolerance': '0', 'key_normalization': {'trim': False, 'case': 'NONE', 'unicode_normalization': 'NONE'}}})
        recon = execute('/recon/runs', {'control_id': control['id'], 'source_version_id': version['id'], 'target_version_id': accepted['id']})
        assert recon['metrics']['matched'] == 110 and recon['metrics']['source_only'] == 6 and recon['metrics']['duplicate_source'] == 4
        monitor = post('/monitors', json={'name': f'Monitor certificado {stamp}', 'dataset_id': accepted['dataset_id'], 'config': fixtures.MONITOR_CONFIG})
        baseline = execute(f"/monitors/{monitor['id']}/runs", {'dataset_version_id': accepted['id']})
        assert baseline['metrics']['failed_checks'] == 0
        rows2 = [dict(r) for r in original[:90]]
        for index, column in enumerate(['customer_id', 'email', 'city']):
            rows2[index][column] = None
        version2 = upload(accepted['dataset_id'], rows2)
        sentinel2 = execute(f"/monitors/{monitor['id']}/runs", {'dataset_version_id': version2['id']})
        assert sentinel2['metrics']['failed_checks'] == 4 and sentinel2['metrics']['health_score'] == 55.56
        rows3 = [{k: v for k, v in r.items() if k != 'source_system'} for r in original[:90]]
        version3 = upload(accepted['dataset_id'], rows3, [c for c in fixtures.COLUMNS if c != 'source_system'])
        sentinel = execute(f"/monitors/{monitor['id']}/runs", {'dataset_version_id': version3['id']})
        assert sentinel['metrics']['failed_checks'] == 1 and sentinel['metrics']['health_score'] == 88.89
        reports = []
        for run in [intake, recon, sentinel]:
            response = client.get(f"/runs/{run['id']}/export.xlsx")
            response.raise_for_status()
            assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            name = f"trackvance_{run['module']}_{run['id']}.xlsx"
            assert name in response.headers['content-disposition']
            workbook = load_workbook(io.BytesIO(response.content))
            assert len(workbook.sheetnames) == 4
            assert all(c.data_type != 'f' for sheet in workbook for row in sheet for c in row)
            target = output / name
            target.write_bytes(response.content)
            reports.append({'module': run['module'], 'run_id': run['id'], 'path': str(target.resolve()), 'artifact_id': response.headers['X-Artifact-ID'], 'status': run['status'], 'decision': run['decision'], 'sheets': workbook.sheetnames})
        result = {'generated_at': stamp, 'base_url': base_url, 'dataset_id': dataset['id'], 'reports': reports, 'certified_sentinel': [{'run_id': r['id'], 'failed_checks': r['metrics']['failed_checks'], 'health_score': r['metrics']['health_score']} for r in [baseline, sentinel2, sentinel]]}
        (output / 'examples.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--base-url', default='http://localhost:3100')
    parser.add_argument('--output', type=Path, default=ROOT / 'outputs/corrections')
    options = parser.parse_args()
    print(json.dumps(generate(options.base_url, options.output), ensure_ascii=False, indent=2))
