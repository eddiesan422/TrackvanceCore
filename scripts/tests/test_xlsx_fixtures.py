"""The volume oracle must detect altered values, not only row counts."""
import hashlib
import importlib.util
import json
import sys
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "xlsx_fixtures.py"
spec = importlib.util.spec_from_file_location("xlsx_fixtures", SCRIPT)
fixtures = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = fixtures
spec.loader.exec_module(fixtures)


@pytest.mark.parametrize("strings", ["inline", "shared"])
def test_fixture_independent_oracle_and_official_reader(tmp_path, strings):
    result = fixtures.generate(tmp_path, 103, strings=strings)
    assert result["oracle"]["status"] == "PASS"
    assert result["oracle"]["null_count"] == 1
    assert result["oracle"]["empty_count"] == 2
    workbook = load_workbook(result["path"], read_only=True, data_only=False)
    try:
        sheet = workbook["Datos"]
        assert sheet.calculate_dimension() == "A1:A1"
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        for _ in range(fixtures.HEADER_ROW):
            header = next(rows)
        assert tuple(header) == fixtures.COLUMNS
        first = next(rows)
        assert first[0] == "000000000001"
        assert first[1] == 2.00000001
        assert first[5].date().isoformat() == "2023-01-02"
        assert first[6] == "=A5+1"
        assert json.loads(Path(result["path"]).with_suffix(".json").read_text(encoding="utf-8"))["rows"] == 103
    finally:
        workbook.close()


def test_oracle_detects_changed_business_value_and_same_count(tmp_path):
    result = fixtures.generate(tmp_path, 3)
    source, changed = Path(result["path"]), tmp_path / "changed.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(changed, "x") as output:
        for info in original.infolist():
            content = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet2.xml":
                content = content.replace(b"000000000003", b"999999999999")
            output.writestr(info, content)
    with pytest.raises(AssertionError, match="Oracle mismatch"):
        fixtures.verify(changed, result)


@pytest.mark.parametrize("strings", ["inline", "shared"])
def test_functional_type_transition_preserves_full_oracle_and_chained_schema(tmp_path, strings):
    from trackvance.acquisition_config import AcquisitionLimits
    from trackvance.batch_readers import FileBatchReader
    from trackvance.dataset_scans import profile_paths
    from trackvance.processing import ProcessingError

    fixture = fixtures.generate(tmp_path, 120, strings=strings, late_type_after=60)
    assert fixture["late_type_after"] == 60 and fixture["oracle"]["status"] == "PASS"
    limits = replace(AcquisitionLimits(), batch_rows=20, xlsx_max_rows=120)
    reader = FileBatchReader(Path(fixture["path"]), Path(fixture["path"]).name, {"sheet_name": "Datos"}, limits)
    paths, digest, count = [], hashlib.sha256(), 0
    for index, batch in enumerate(reader):
        assert batch.frame.height <= 20
        for row in batch.frame.iter_rows(named=True):
            count += 1
            assert row == fixtures.expected_row(count, late_type_after=60)
            digest.update(fixtures.canonical(row))
        path = tmp_path / f"accepted-{index}.parquet"
        batch.physical_frame().write_parquet(path)
        paths.append(path)
    assert count == 120 and digest.hexdigest() == fixture["canonical_rows_sha256"]
    assert reader.native_schema["late_type"] == "String"
    acquired, _, _ = profile_paths(paths, native_types=reader.native_schema, limits=limits)
    accepted, _, _ = profile_paths(paths, limits=limits)
    acquired_types = {column["name"]: column["logical_type"] for column in acquired}
    accepted_types = {column["name"]: column["logical_type"] for column in accepted}
    assert acquired_types == accepted_types
    assert accepted_types["late_type"] == "STRING"
    assert accepted_types["amount"] == "DECIMAL" and accepted_types["date"] == "DATE"
    excess = fixtures.generate(tmp_path, 121, strings=strings, late_type_after=60)
    assert excess["late_type_after"] == 60 and excess["oracle"]["status"] == "PASS"
    with pytest.raises(ProcessingError, match="ACQUISITION_ROW_LIMIT"):
        list(FileBatchReader(Path(excess["path"]), Path(excess["path"]).name, {"sheet_name": "Datos"}, limits))


def test_default_type_transition_stays_at_historical_hundred_thousand_rows(tmp_path):
    fixture = fixtures.generate(tmp_path, 120)
    assert "late_type_after" not in fixture
    assert Path(fixture["path"]).name == "xlsx-120-inline-diverse.xlsx"
    assert fixtures.expected_row(100000)["late_type"] == str(100000 % 997)
    assert fixtures.expected_row(100001)["late_type"] == "late-text-100001"
