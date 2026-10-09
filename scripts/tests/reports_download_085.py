"""R085-01 real proxy downloads, with complete independent client-side oracles.

This host runner never imports application services or reads server output files.
CSV/XLSX results are streamed to its own private client directory. A low-volume
functional run uses --rows 120 --limit 120; 400k/1m belong exclusively to local deep.
Docker startup is optional for an existing guarded context. No images are built.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import http.client
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import zipfile
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import Path, PurePosixPath
from uuid import uuid4
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "scripts/tests"))
import catalog_reports_api_cycle as cycle
import certification_v080 as guard
from browser_evidence import run_browser
from ci.host_resources import limit_cpu_affinity, sample_process_tree, set_cpu_affinity
from ci.owned_cleanup import cleanup as cleanup_owned
from ci.owned_cleanup import snapshot as docker_snapshot
from ci_images import verified_images
from report_resources_085 import configure as configure_resources
from report_resources_085 import validate as validate_resources

HEADERS = ["transaction_id", "customer_id", "customer_name", "tx_num", "amount", "day", "instant", "active", "optional", "formula", "unicode"]
LOGICAL_TYPES = ["STRING", "STRING", "STRING", "INT64", "DECIMAL", "DATE", "TIMESTAMP", "BOOLEAN", "STRING", "STRING", "STRING"]
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
SCOPES = ("postgres", "api", "worker", "acquisition-worker", "web")


class CheckFailed(ValueError):
    """Closed diagnostic: fixture values, cookies and credentials never enter evidence."""


def require(condition, code):
    if not condition:
        raise CheckFailed(code)


def write_json(path: Path, value: dict):
    temporary = path.with_suffix(path.suffix + ".partial")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while block := stream.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


class Resources:
    def __init__(self, max_bytes: int, deadline: float):
        self.max_bytes, self.deadline = max_bytes, deadline
        self.stop = threading.Event()
        self.peak, self.samples, self.available = 0, 0, False
        self.thread = threading.Thread(target=self._sample, daemon=True)

    def _sample(self):
        while not self.stop.is_set():
            snapshot = sample_process_tree(os.getpid())
            if snapshot.get("status") == "PASS":
                self.available = True
                self.peak = max(self.peak, sum(row["rss_bytes"] for row in snapshot["processes"]))
                self.samples += 1
            self.stop.wait(0.25)

    def __enter__(self):
        self.thread.start()
        return self

    def check(self):
        require(time.monotonic() < self.deadline, "CLIENT_TOTAL_DEADLINE")
        require(self.peak <= self.max_bytes, "CLIENT_MEMORY_BUDGET")

    def __exit__(self, *_):
        self.stop.set()
        self.thread.join(timeout=2)

    def metrics(self):
        return {"scope": "HOST_CLIENT_PROCESS_TREE", "status": "PASS" if self.available else "UNAVAILABLE",
                "sampled_peak_rss_bytes": self.peak, "sampling_interval_seconds": 0.25,
                "samples": self.samples, "budget_bytes": self.max_bytes}


class ProxyClient(cycle.Client):
    def __init__(self, port: int, timeout: int):
        super().__init__()
        self.port, self.timeout = port, timeout
        self.base = f"http://127.0.0.1:{port}/api/v1"

    def upload(self, path: Path):
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=self.timeout)
        try:
            connection.putrequest("POST", "/api/v1/datasets/uploads/stage?filename=" + urllib.parse.quote(path.name))
            for name, value in {"Content-Type": "application/octet-stream", "Content-Length": str(path.stat().st_size),
                                "X-CSRF-Token": self.csrf,
                                "Cookie": "; ".join(f"{c.name}={c.value}" for c in self.cookies)}.items():
                connection.putheader(name, value)
            connection.endheaders()
            with path.open("rb") as source:
                while chunk := source.read(65536):
                    connection.send(chunk)
            response = connection.getresponse()
            value = json.loads(response.read())
            require(response.status == 201, "SYNTHETIC_UPLOAD_FAILED")
            return value["upload"]["id"]
        finally:
            connection.close()


def terminal(client: ProxyClient, path: str, resources: Resources, timeout: int = 1800):
    end = min(resources.deadline, time.monotonic() + timeout)
    while time.monotonic() < end:
        resources.check()
        value = client.call(path)
        if value["status"] not in {"QUEUED", "RUNNING", "PENDING"}:
            return value
        time.sleep(0.5)
    raise CheckFailed("TERMINAL_STATE_TIMEOUT")


def customer(index: int):
    return [f"{index:08d}", f'Cliente ñ 東京😀 {index}, "literal"']


def transaction(index: int, customers: int, *, canonical: bool = False):
    number = index + 1
    instant = datetime(2026, 1, 1, 0, 0, 0, 123456, tzinfo=UTC) + timedelta(seconds=index)
    source_instant = instant.isoformat().replace("+00:00", "Z")
    # Offset crosses calendar boundaries and preserves the same exact instant.
    if index % 2:
        from datetime import timezone
        source_instant = instant.astimezone(timezone(timedelta(hours=-5))).isoformat()
    optional = [None, "", r"\N", "'literal", " "][index % 5]
    return [f"{number:012d}", customer(index % customers)[0], number,
            str(Decimal("12345678901234567890.12345678") + Decimal(number) / Decimal(100000000)),
            (date(2026, 1, 1) + timedelta(days=index % 28)).isoformat(),
            instant.isoformat() if canonical else source_instant,
            bool(index % 2) if canonical else "true" if index % 2 else "false",
            optional, "=SUM(1,2)" if index % 2 else "-001", 'á東京😀,%_ "comillas"\\\nsegunda línea']


def expected_rows(count: int, customers: int):
    for index in range(count):
        values = transaction(index, customers, canonical=True)
        yield [*values[:2], customer(index % customers)[1], *values[2:]]


def population_hash(rows, resources: Resources | None = None):
    digest, count, logical_bytes = hashlib.sha256(), 0, 0
    for row in rows:
        serialized = json.dumps(row, ensure_ascii=False, separators=(",", ":")).encode() + b"\n"
        digest.update(serialized)
        logical_bytes += len(serialized)
        count += 1
        if resources and count % 1000 == 0:
            resources.check()
    return {"rows": count, "ordered_logical_sha256": digest.hexdigest(), "logical_json_bytes": logical_bytes}


def fixture(path: Path, count: int, customers: int, kind: str, resources: Resources):
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output, quoting=csv.QUOTE_ALL)
        writer.writerow(["customer_id", "customer_name"] if kind == "customers" else [name for name in HEADERS if name != "customer_name"])
        for index in range(count):
            row = customer(index) if kind == "customers" else transaction(index, customers)
            # Quote empty strings; only null is unquoted empty under the upload contract.
            parts = []
            for value in row:
                if value is None:
                    parts.append("")
                else:
                    text = str(value)
                    parts.append('"' + text.replace('"', '""') + '"')
            output.write(",".join(parts) + "\r\n")
            if index % 1000 == 0:
                resources.check()


def acquire(client: ProxyClient, path: Path, alias: str, count: int, macro: str, domain: str, resources: Resources):
    began = time.monotonic()
    upload = client.upload(path)
    dataset = client.call("/datasets", {"name": "Descarga 085 " + alias + " " + os.urandom(6).hex()}, expected=201)
    names = ["customer_id", "customer_name"] if alias == "c" else [name for name in HEADERS if name != "customer_name"]
    overrides = {name: {"logical_type": LOGICAL_TYPES[HEADERS.index(name)],
                        "semantic_tag": "IDENTIFIER" if name in {"transaction_id", "customer_id"} else None} for name in names}
    acquisition = client.call(f"/datasets/{dataset['id']}/acquisitions", {"upload_id": upload, "column_overrides": overrides}, expected=202)
    acquired = terminal(client, "/acquisitions/" + acquisition["id"], resources)
    require(acquired["status"] == "SUCCESS" and acquired["processed_rows"] == count, "ACQUISITION_POPULATION")
    client.call(f"/datasets/{dataset['id']}/governance", {"expected_version": 1, "macro_domain_id": macro,
                "domain_id": domain, "information_classification": "INTERNAL"}, method="PATCH")
    contract = client.call("/intake/contracts", {"name": "Descarga 085 control " + alias + " " + os.urandom(6).hex(),
                           "dataset_id": dataset["id"], "config": {"required_columns": ["customer_id"], "max_error_rate": 0}}, expected=201)
    run = client.call("/intake/runs", {"contract_id": contract["id"], "dataset_version_id": acquired["output_version_id"], "requested_engine": "AUTO"}, expected=202)
    approved = terminal(client, "/runs/" + run["id"], resources)
    require(approved["status"] == "SUCCESS" and approved["decision"] == "APPROVED", "INTAKE_STRICT_DECISION")
    require(approved["metrics"]["validation_coverage_rows"] == approved["metrics"]["total_rows"] == count, "INTAKE_STRICT_COVERAGE")
    schema = client.call("/dataset-versions/" + approved["output_version_id"] + "/profile")["schema"]
    require({c["name"]: c["logical_type"] for c in schema} == {n: overrides[n]["logical_type"] for n in names}, "INTAKE_SCHEMA_PRESERVED")
    require({c["name"]: c.get("semantic_tag") for c in schema} == {n: overrides[n]["semantic_tag"] for n in names}, "INTAKE_IDENTIFIER_TAGS_PRESERVED")
    return {"alias": alias, "input_dataset_id": dataset["id"], "contract_id": contract["id"],
            "contract_revision_ids": [contract["id"]], "policy": "LATEST_APPROVED"}, {
                "alias": alias, "rows": count, "acquisition_id": acquired["id"], "approval_id": approved["id"],
                "output_version_id": approved["output_version_id"], "schema": schema,
                "engine": approved.get("execution_plan", {}).get("engine"),
                "input_fixture_bytes": path.stat().st_size, "input_fixture_sha256": file_hash(path),
                "duration_seconds": round(time.monotonic() - began, 3), "status": "PASS"}


def query(sources: list[dict], count: int, *, variant: str = "WHERE"):
    selected = ",".join(("c." if name == "customer_name" else "t.") + name + " AS " + name for name in HEADERS)
    sql = f"SELECT {selected} FROM t INNER JOIN c ON t.customer_id=c.customer_id"
    parameters = [{"name": "final_rows", "type": "INTEGER", "value": str(count)}]
    if variant == "WHERE":
        sql += " WHERE t.tx_num <= $final_rows"
    elif variant == "LIMIT":
        sql += f" ORDER BY t.tx_num LIMIT {count}"
        parameters = []
    else:
        raise CheckFailed("QUERY_VARIANT")
    if variant != "LIMIT":
        sql += " ORDER BY t.tx_num"
    return {"mode": "SQL", "sources": sources, "sql": sql, "parameters": parameters,
            "joins": [{"left_alias": "t", "right_alias": "c", "type": "INNER",
                       "keys": [{"left_column": "customer_id", "right_column": "customer_id"}],
                       "expected_cardinality": "N:1", "allow_many_to_many": False}],
            "columns": [], "source_filters": {}, "order_by": [], "expected_schemas": {}}


def csv_population(path: Path):
    with path.open("r", encoding="utf-8", newline="") as source:
        reader = csv.reader(source)
        require(next(reader) == HEADERS, "CSV_HEADERS")
        for row in reader:
            require(len(row) == len(HEADERS), "CSV_COLUMN_COUNT")
            values = [cycle.decode_csv(value) for value in row]
            values[3] = int(values[3])
            require(values[7] in {"true", "false"}, "CSV_BOOLEAN")
            values[7] = values[7] == "true"
            yield values


def column_index(reference: str) -> int:
    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", reference)
    require(match, "XLSX_CELL_REFERENCE")
    result = 0
    for letter in match[1]:
        result = result * 26 + ord(letter) - ord("A") + 1
    return result - 1


def xlsx_population(path: Path):
    with zipfile.ZipFile(path) as package, package.open("xl/worksheets/sheet1.xml") as stream:
        parent, count = None, 0
        for event, element in ET.iterparse(stream, events=("start", "end")):
            if event == "start" and element.tag == NS + "sheetData":
                parent = element
            if event != "end" or element.tag != NS + "row":
                continue
            count += 1
            require(element.get("r") == str(count), "XLSX_PHYSICAL_ROW_SEQUENCE")
            values = [None] * len(HEADERS)
            seen = set()
            for cell in element:
                index = column_index(cell.get("r", ""))
                require(index < len(HEADERS) and index not in seen and re.search(r"[0-9]+$", cell.get("r", ""))[0] == str(count), "XLSX_CELL_POSITION")
                seen.add(index)
                require(cell.find(NS + "f") is None, "XLSX_FORMULA_CREATED")
                kind = cell.get("t")
                if kind == "inlineStr":
                    text = cell.find(NS + "is/" + NS + "t")
                    require(text is not None and text.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve", "XLSX_INLINE_TEXT")
                    values[index] = text.text or ""
                elif kind == "b":
                    require(cell.findtext(NS + "v") in {"0", "1"}, "XLSX_BOOLEAN")
                    values[index] = cell.findtext(NS + "v") == "1"
                elif kind == "n":
                    values[index] = int(cell.findtext(NS + "v"))
                else:
                    raise CheckFailed("XLSX_CELL_TYPE")
            if count == 1:
                require(values == HEADERS, "XLSX_HEADERS")
            else:
                require(type(values[3]) is int and type(values[7]) is bool, "XLSX_TYPED_CELLS")
                for index in (0, 1, 2, 4, 5, 6, 9, 10):
                    require(isinstance(values[index], str), "XLSX_EXACT_TEXT_CELLS")
                yield values
            element.clear()
            require(parent is not None, "XLSX_SHEET_DATA")
            parent.remove(element)  # Clearing alone leaves one empty node per row.


def validate_package(path: Path):
    with zipfile.ZipFile(path) as package:
        entries = package.infolist()
        names = [entry.filename for entry in entries]
        required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml"}
        require(set(names) == required and len(names) == len(set(names)), "XLSX_PACKAGE_ENTRIES")
        require(package.testzip() is None, "XLSX_ZIP_CRC")
        for name in required - {"xl/worksheets/sheet1.xml"}:
            require(package.getinfo(name).file_size < 1024 * 1024, "XLSX_METADATA_BOUND")
        content = ET.fromstring(package.read("[Content_Types].xml"))
        require(content.tag == CT + "Types", "XLSX_CONTENT_NAMESPACE")
        overrides = {node.get("PartName"): node.get("ContentType") for node in content if node.tag == CT + "Override"}
        require(overrides.get("/xl/workbook.xml", "").endswith("spreadsheetml.sheet.main+xml") and overrides.get("/xl/worksheets/sheet1.xml", "").endswith("spreadsheetml.worksheet+xml"), "XLSX_CONTENT_TYPES")
        rootrels = ET.fromstring(package.read("_rels/.rels"))
        require(rootrels.tag == REL + "Relationships" and any(node.get("Target") == "xl/workbook.xml" and node.get("Type", "").endswith("/officeDocument") for node in rootrels), "XLSX_ROOT_RELATIONSHIP")
        workbook = ET.fromstring(package.read("xl/workbook.xml"))
        require(workbook.tag == NS + "workbook", "XLSX_WORKBOOK_NAMESPACE")
        sheets = workbook.findall(NS + "sheets/" + NS + "sheet")
        require(len(sheets) == 1, "XLSX_WORKSHEET_COUNT")
        relations = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        identifier = sheets[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        require(relations.tag == REL + "Relationships" and any(node.get("Id") == identifier and node.get("Type", "").endswith("/worksheet") and str(PurePosixPath("xl") / node.get("Target", "")) == "xl/worksheets/sheet1.xml" for node in relations), "XLSX_WORKSHEET_RELATIONSHIP")
        return {"status": "PASS", "crc": "ALL_ENTRIES_PASS", "expanded_bytes": sum(entry.file_size for entry in entries),
                "serialized_zip_bytes": path.stat().st_size, "zip64_present": any(entry.extract_version >= 45 for entry in entries),
                "entries": len(entries), "namespaces_content_types_relationships": "PASS"}


def openpyxl_population(path: Path):
    from openpyxl import load_workbook
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False)
        try:
            require(len(workbook.worksheets) == 1, "READER_WORKSHEET_COUNT")
            iterator = workbook.active.iter_rows(min_col=1, max_col=len(HEADERS))
            require([cell.value for cell in next(iterator)] == HEADERS, "READER_HEADERS")
            for cells in iterator:
                require(all(cell.data_type != "f" for cell in cells), "READER_FORMULA_CREATED")
                yield [cell.value for cell in cells]
        finally:
            workbook.close()


def download(client: ProxyClient, context_id: str, format_name: str, path: Path, expected: dict, resources: Resources):
    began, identity, bytes_received = time.monotonic(), None, 0
    partial = path.with_suffix(path.suffix + ".incomplete")
    try:
        with client.stream(context_id, format_name) as response, partial.open("xb") as output:
            identity = response.headers.get("X-Report-Execution-Id")
            require(identity and response.status == 200 and response.headers.get("Cache-Control") == "no-store", "DOWNLOAD_IDENTITY_HEADERS")
            while block := response.read(65536):
                resources.check()
                output.write(block)
                bytes_received += len(block)
        transfer_seconds = round(time.monotonic() - began, 3)
        state = terminal(client, "/reports/executions/" + identity, resources, timeout=30)
        require(state["status"] == "SUCCESS" and state["generation_status"] == "COMPLETE" and state["transmission_status"] == "COMPLETE", "DOWNLOAD_NOT_SUCCESS")
        require(state["metrics"].get("rows") == expected["rows"] and state["metrics"].get("serialized_bytes") == bytes_received
                and state["metrics"].get("serialization_complete") is True, "DOWNLOAD_TERMINAL_COUNTS")
    except BaseException:  # noqa: TRY203 - document the ownership of incomplete client files.
        # Remains explicitly .incomplete and is never included among valid outputs.
        raise
    verification_started = time.monotonic()
    result = {"format": format_name, "execution_id": identity, "bytes_received": bytes_received,
              "transfer_seconds": transfer_seconds, "server_metrics": state["metrics"], "terminal_status": state["status"],
              "generation_status": state["generation_status"], "transmission_status": state["transmission_status"],
              "client_file": path.name, "file_sha256": file_hash(partial)}
    if format_name == "CSV":
        result["independent_csv"] = population_hash(csv_population(partial), resources)
        require(result["independent_csv"] == expected, "CSV_COMPLETE_POPULATION_ORACLE")
    else:
        result["ooxml"] = validate_package(partial)
        result["independent_xml"] = population_hash(xlsx_population(partial), resources)
        require(result["independent_xml"] == expected, "XLSX_COMPLETE_XML_POPULATION_ORACLE")
        result["independent_openpyxl"] = population_hash(openpyxl_population(partial), resources)
        require(result["independent_openpyxl"] == expected, "XLSX_COMPLETE_READER_POPULATION_ORACLE")
        result["physical_rows"] = expected["rows"] + 1
        result["excel_open_without_repair"] = "NOT_EXECUTED_NO_ISOLATED_EXCEL_SESSION"
    result.update(status="PASS", verification_seconds=round(time.monotonic() - verification_started, 3))
    partial.replace(path)
    return result


def rejected_download(client: ProxyClient, context_id: str, format_name: str, resources: Resources):
    try:
        with client.stream(context_id, format_name) as response:
            response.read(1)
        raise CheckFailed("EXCESS_RECEIVED_OUTPUT_BYTES")
    except urllib.error.HTTPError as error:
        value = json.load(error)
        require(error.code == 422 and value.get("error", {}).get("code") == "REPORT_RESULT_LIMIT", "EXCESS_NOT_REJECTED_FINAL")
        identity = value["error"].get("details", {}).get("execution_id")
        require(identity, "EXCESS_EXECUTION_IDENTITY")
        state = terminal(client, "/reports/executions/" + identity, resources, timeout=30)
        require(state["status"] == "FAILED" and state["error_code"] == "REPORT_RESULT_LIMIT"
                and state["generation_status"] == "FAILED" and state["transmission_status"] == "NOT_STARTED"
                and state.get("output_version_id") is None, "EXCESS_TERMINAL_STATE")
        return {"status": "PASS", "format": format_name, "execution_id": identity, "http_status": error.code,
                "result_bytes_received": 0, "error_code": state["error_code"], "terminal_status": state["status"],
                "generation_status": state["generation_status"], "transmission_status": state["transmission_status"],
                "output_version_id": state.get("output_version_id")}


def browser_download(client: ProxyClient, context: dict, sources: list[dict], summary: dict,
                     client_dir: Path, size: int, resources: Resources):
    """The UI writes to a real OPFS handle; the host verifies its complete file."""
    pnpm = shutil.which("pnpm")
    require(pnpm, "BROWSER_RUNNER_UNAVAILABLE")
    draft = query(sources, size)
    draft["expected_schemas"] = {source["alias"]: source["schema"] for source in summary["sources"]}
    draft["output_preferences"] = {"format": "XLSX"}
    definition = client.call("/reports/definitions", {"name": "OPFS085 " + summary["attempt_id"], "draft": draft})
    fixture_path = client_dir / "browser-fixture.json"
    write_json(fixture_path, {"definition_id": definition["id"], "rows": size,
               "project": context["project"], "source_sha": summary["source_sha"],
               "client_directory": str(client_dir), "filename": "browser-result.xlsx.incomplete"})
    counters = run_browser(pnpm, ["tests-e2e/reports-download-opfs.spec.ts"], root=ROOT,
        project=context["project"] + "-opfs", evidence=client_dir,
        environment={**os.environ, "TV_E2E_URL": f"http://localhost:{context['port']}",
                     "TV_REPORT_DOWNLOAD_OPFS": "true", "TV_DOWNLOAD_085_FIXTURE": str(fixture_path)},
        timeout_seconds=min(1800, int(resources.deadline - time.monotonic())), resource_check=resources.check)
    require(counters.get("status") == "PASS" and counters.get("expected") == 1
            and all(counters.get(key, 0) == 0 for key in ("skipped", "unexpected", "flaky")), "BROWSER_OPFS_NOT_COMPLETED")
    receipt = json.loads((client_dir / "browser-receipt.json").read_text(encoding="utf-8"))
    require(receipt.get("status") == "PASS" and receipt.get("method") == "REAL_OPFS_FILE_HANDLE"
            and receipt.get("rows") == size and receipt.get("source_sha") == summary["source_sha"], "BROWSER_OPFS_RECEIPT")
    path = client_dir / "browser-result.xlsx.incomplete"
    require(path.is_file() and path.stat().st_size == receipt.get("bytes_received"), "BROWSER_OPFS_CLIENT_FILE")
    state = terminal(client, "/reports/executions/" + receipt["execution_id"], resources, timeout=30)
    require(state["status"] == "SUCCESS" and state["generation_status"] == "COMPLETE"
            and state["transmission_status"] == "COMPLETE" and state["metrics"].get("serialized_bytes") == path.stat().st_size
            and state["metrics"].get("rows") == size, "BROWSER_OPFS_SERVER_TERMINAL")
    package = validate_package(path)
    xml = population_hash(xlsx_population(path), resources)
    reader = population_hash(openpyxl_population(path), resources)
    require(xml == reader == summary["oracles"][str(size)], "BROWSER_OPFS_COMPLETE_ORACLE")
    digest = file_hash(path)
    final = client_dir / "browser-result.xlsx"
    path.replace(final)
    return {"status": "PASS", **receipt, "browser": counters, "client_file": final.name,
            "file_sha256": digest, "independent_xml": xml, "independent_openpyxl": reader,
            "ooxml": package, "physical_rows": size + 1, "server_metrics": state["metrics"],
            "generation_status": state["generation_status"], "transmission_status": state["transmission_status"],
            "excel_open_without_repair": "NOT_EXECUTED_NO_ISOLATED_EXCEL_SESSION"}


def certify(directory: Path, context: dict, sizes: list[int], limit: int, source_sha: str, resources: Resources,
            *, with_browser: bool = False):
    attempt = uuid4().hex
    client_dir = directory / ("client-download-085-" + attempt[:12])
    client_dir.mkdir(exist_ok=False)
    evidence = directory / "reports-download-085.json"
    archive = directory / ("reports-download-085-" + attempt + ".json")
    summary = {"requirements": ["R080-02", "R080-04", "R085-01", "R085-03"], "implementation": "0.8.5",
               "source_sha": source_sha, "attempt_id": attempt, "project": context["project"], "status": "RUNNING", "downloads": [],
               "sources": [], "client_result_storage": "PRIVATE_HOST_ONLY", "result_columns": HEADERS,
               "concurrency": 1, "requested_result_rows": sizes, "product_row_limit_under_test": limit,
               "server_filesystem_observation": "SEPARATE_TRACE_GATE_REQUIRED", "output_data_published": False}
    summary["browser_streaming_requested"] = with_browser
    def phase(value):
        resources.check()
        summary["active_phase"] = value
        summary["client_resources"] = resources.metrics()
        write_json(evidence, summary)
        write_json(archive, summary)
        print(json.dumps({"phase": value, "project": context["project"]}), flush=True)
    try:
        phase("AUTHENTICATION")
        client = ProxyClient(context["port"], 1900)
        client.call("/auth/demo", {})
        limits = client.call("/reports/limits")
        require(limits["profiles"]["DOWNLOAD"]["max_rows"] == limits["profiles"]["XLSX"]["max_rows"] == limit, "EFFECTIVE_ROW_LIMIT")
        require(limits["profiles"]["PREVIEW"]["max_rows"] == 10, "PREVIEW_TEN_ROWS")
        summary["effective_limits"] = limits
        macro = client.call("/catalog/macrodomains", {"name": "Descarga085 " + os.urandom(6).hex()}, expected=201)
        domain = client.call("/catalog/domains", {"name": "Descargas", "macro_domain_id": macro["id"]}, expected=201)
        customers, population = min(100000, max(1, limit // 10)), limit + 1
        summary.update(customers=customers, transactions_input=population)
        sources = []
        for alias, kind, count in [("t", "transactions", population), ("c", "customers", customers)]:
            phase("FIXTURE_" + alias.upper())
            path = client_dir / (kind + ".csv")
            fixture(path, count, customers, kind, resources)
            phase("ACQUIRE_INTAKE_" + alias.upper())
            source, proof = acquire(client, path, alias, count, macro["id"], domain["id"], resources)
            sources.append(source)
            summary["sources"].append(proof)
        for size in sizes:
            phase("RESOLVE_" + str(size))
            frozen = client.call("/reports/resolve", {"draft": query(sources, size)})
            summary.setdefault("contexts", []).append({"rows": size, "context_id": frozen["context_id"], "query_hash": frozen["query_hash"], "sources": frozen["sources"]})
            preview = client.call("/reports/preview", {"context_id": frozen["context_id"]})
            require(len(preview["rows"]) == min(10, size) and preview["status"] == "SUCCESS", "PREVIEW_RESULT")
            require([[row[name] for name in HEADERS] for row in preview["rows"]] == list(expected_rows(min(10, size), customers)), "PREVIEW_VALUE_ORACLE")
            phase("ORACLE_" + str(size))
            expected = population_hash(expected_rows(size, customers), resources)
            summary.setdefault("oracles", {})[str(size)] = expected
            for format_name in ("CSV", "XLSX"):
                phase("DOWNLOAD_" + format_name + "_" + str(size))
                summary["downloads"].append(download(client, frozen["context_id"], format_name,
                    client_dir / f"result-{size}.{format_name.lower()}", expected, resources))
        phase("EXCESS_FINAL_RESULT")
        excess = client.call("/reports/resolve", {"draft": query(sources, limit + 1)})
        summary["excess"] = [rejected_download(client, excess["context_id"], format_name, resources) for format_name in ("CSV", "XLSX")]
        # This query's join has limit+1 rows but FINAL LIMIT is admitted and exact.
        phase("FINAL_LIMIT_AFTER_JOIN")
        bounded = client.call("/reports/resolve", {"draft": query(sources, sizes[0], variant="LIMIT")})
        summary["final_limit"] = download(client, bounded["context_id"], "CSV", client_dir / "final-limit.csv",
                                          summary["oracles"][str(sizes[0])], resources)
        if with_browser:
            phase("REAL_BROWSER_OPFS_DOWNLOAD")
            summary["browser_streaming"] = browser_download(client, context, sources, summary, client_dir, max(sizes), resources)
        summary.update(status="PASS", active_phase="COMPLETE", client_resources=resources.metrics())
        require(resources.available, "CLIENT_RESOURCE_SAMPLING_UNAVAILABLE")
        write_json(evidence, summary)
        write_json(archive, summary)
        return summary
    except BaseException as error:
        summary.update(status="FAIL", error_type=type(error).__name__, error_code=str(error) if isinstance(error, CheckFailed) else "CERTIFICATION_EXCEPTION",
                       client_resources=resources.metrics())
        write_json(evidence, summary)
        write_json(archive, summary)
        raise


def prepare(port: int, main_project: str | None, limit: int):
    require(verified_images() is not None, "SAME_SHA_VERIFIED_IMAGES_REQUIRED")
    directory = guard.init("reports-download085", port, main_project)
    directory, context = guard.load_context(directory)
    configure_resources(directory, context, limit)
    guard.preflight(directory, context)
    return directory, context


def verify_code(source_sha: str):
    require(re.fullmatch(r"[a-f0-9]{40}", source_sha), "SOURCE_SHA_REQUIRED")
    actual = guard.command(["git", "rev-parse", "HEAD"]).strip()
    require(actual == source_sha, "SOURCE_SHA_MISMATCH")
    require(not guard.command(["git", "status", "--porcelain", "--", "backend", "frontend", "scripts", "compose.yml", "deploy"]).strip(), "EXECUTABLE_WORKTREE_NOT_COMMITTED")
    images = verified_images()
    require(images is not None, "SAME_SHA_VERIFIED_IMAGES_REQUIRED")
    return images


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", nargs="+", type=int, choices=(120, 400000, 1000000), default=[400000, 1000000])
    parser.add_argument("--limit", type=int, choices=(120, 1000000), default=1000000)
    parser.add_argument("--context", type=Path)
    parser.add_argument("--port", type=int, default=32085)
    parser.add_argument("--main-project")
    parser.add_argument("--source-sha", default=os.getenv("TRACKVANCE_SOURCE_SHA", os.getenv("CI_SOURCE_SHA", "")))
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--start", action="store_true", help="Start absent owned services of an existing context")
    parser.add_argument("--keep", action="store_true", help="Retain new owned resources for another explicitly planned test")
    parser.add_argument("--with-browser", action="store_true", help="Require real UI streaming to an OPFS FileSystemFileHandle")
    parser.add_argument("--client-memory-mib", type=int)
    parser.add_argument("--client-cpus", type=int, default=1)
    parser.add_argument("--timeout", type=int, default=7200)
    args = parser.parse_args()
    if args.client_memory_mib is None:
        args.client_memory_mib = 1024 if args.with_browser else 512
    require(all(value <= args.limit for value in args.rows) and len(set(args.rows)) == len(args.rows), "RESULT_SIZES")
    require(not any(value > 120 for value in args.rows) or args.with_browser, "DEEP_REAL_BROWSER_REQUIRED")
    require(not os.getenv("GITHUB_ACTIONS") or args.rows == [120] and args.limit == 120, "DEEP_VOLUME_FORBIDDEN_IN_ACTIONS")
    require(128 <= args.client_memory_mib <= 2048 and 1 <= args.client_cpus <= 4 and 60 <= args.timeout <= 10800, "FINITE_CLIENT_BUDGET")
    images = verify_code(args.source_sha)
    before = docker_snapshot()
    directory, context = guard.load_context(args.context) if args.context else prepare(args.port, args.main_project, args.limit)
    resource_profile = validate_resources(json.loads((directory / "compose.json").read_text(encoding="utf-8")))
    require(context["image"] == images["backend"], "CONTEXT_IMAGE_SHA_MISMATCH")
    require(context["project"].startswith("trackvance-v080-test-") and context["project"] != context["main_project"], "OWNED_CONTEXT_REQUIRED")
    if args.prepare_only:
        print(json.dumps({"status": "PREPARED", "context": str(directory), "project": context["project"]}))
        return
    summary, began, affinity = {"status": "FAIL"}, time.monotonic(), None
    try:
        guard.preflight(directory, context)
        affinity = limit_cpu_affinity(args.client_cpus)
        if not args.context or args.start:
            guard.command([*guard.compose_args(directory, context), "up", "--no-build", "--detach", "--wait", "--wait-timeout", "240", *SCOPES], timeout=300)
        with Resources(args.client_memory_mib * 1024**2, time.monotonic() + args.timeout) as resources:
            summary = certify(directory, context, args.rows, args.limit, args.source_sha, resources, with_browser=args.with_browser)
    finally:
        summary.update(duration_seconds=round(time.monotonic() - began, 3), client_affinity=affinity,
                       private_resource_profile=resource_profile,
                       main_unchanged=guard.inventory(context["main_project"]) == context["main_before"])
        if not args.keep:
            try:
                summary["cleanup"] = cleanup_owned(before, projects={context["project"]})
            except (OSError, ValueError, RuntimeError, subprocess.SubprocessError) as error:
                summary.update(status="FAIL", cleanup={"status": "FAIL", "error_type": type(error).__name__})
        else:
            summary["cleanup"] = {"status": "RETAINED_FOR_NEXT_OWNED_TEST", "project": context["project"], "expires_within_seconds": 7200}
        if affinity:
            set_cpu_affinity(affinity["original_logical_processors"])
        summary["main_unchanged"] = guard.inventory(context["main_project"]) == context["main_before"]
        require(summary["main_unchanged"], "PROTECTED_MAIN_CHANGED")
        evidence = directory / "reports-download-085.json"
        # Keep detailed incremental failure evidence written by certify.
        if summary["status"] == "FAIL" and evidence.exists():
            retained = json.loads(evidence.read_text(encoding="utf-8"))
            if retained.get("status") == "FAIL":
                summary = {**retained, **summary}
        write_json(evidence, summary)
        if summary.get("attempt_id"):
            write_json(directory / ("reports-download-085-" + summary["attempt_id"] + ".json"), summary)
        print(json.dumps({"status": summary["status"], "evidence": str(evidence), "main_unchanged": summary["main_unchanged"]}), flush=True)
    require(summary["status"] == "PASS", "DOWNLOAD_CERTIFICATION_FAILED")


if __name__ == "__main__":
    main()
